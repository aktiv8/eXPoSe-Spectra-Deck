"""Word (.docx) export of an experiment: an editable document built from the
same data as the PDF report and the PowerPoint deck, for anyone who needs to
add material afterwards that the app itself does not capture.

Sections, in the order of the ``reportspec`` spec: a cover (letterhead logo,
title, customer / reference / operator / date), the contents, the summary,
the quantification (``resultspages``), the saved figures, the camera
pictures and SnapMaps (``imagepages``), methods and calibration texts, the
acquisition metadata (``metasummary``) and the data files list.

Unlike the PDF, Word lays out its own pages, so this module is a good deal
simpler than ``report.py``: headings use the built-in Word heading styles
(so they show in the Navigation Pane and the reader can restructure the
document freely) and the contents is a native Word TOC field the reader
updates themselves (right-click it, "Update Field") rather than page numbers
computed here. Needs python-docx; no Tk here.
"""

from __future__ import annotations

import datetime
import io
import os

import appinfo
import covers
import metasummary
import reportspec
import resultspages

FONT = "Calibri"
PAGE_W_IN = 6.5           # content width inside the default 1" margins


class DocxError(Exception):
    """The document could not be built (message is user-facing)."""


def _modules():
    try:
        import docx
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError:
        raise DocxError("Word export needs python-docx "
                        "(pip install python-docx).")
    return docx, Inches, Pt, RGBColor, WD_ALIGN_PARAGRAPH, OxmlElement, qn


# -- pure helpers --------------------------------------------------------------
def _paragraphs(text):
    """Blank-line separated paragraphs of free text (single newlines kept)."""
    return [b.strip("\n") for b in (text or "").replace("\r\n", "\n")
            .split("\n\n") if b.strip()]


def _size(n):
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024.0


# -- the document wrapper -------------------------------------------------------
class _Doc:
    def __init__(self, accent=""):
        (self.docx, self.Inches, self.Pt, self.RGB, self.ALIGN, self.Oxml,
         self.qn) = _modules()
        self.doc = self.docx.Document()
        self.accent = covers.valid_accent(accent) or covers.DEFAULT_ACCENT
        self.ink = covers.mix(self.accent, "#000000", 0.2)
        self.alt = covers.tint(self.accent, 0.92)
        self._style_headings()

    def _rgb(self, hex_colour):
        h = hex_colour.lstrip("#")
        return self.RGB(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

    def _style_headings(self):
        ink = self._rgb(self.ink)
        for name, size in (("Title", 26), ("Heading 1", 16),
                           ("Heading 2", 13)):
            try:
                st = self.doc.styles[name]
            except KeyError:
                continue
            st.font.name = FONT
            st.font.size = self.Pt(size)
            st.font.color.rgb = ink

    # -- text --------------------------------------------------------------
    def heading(self, text, level=1):
        return self.doc.add_heading(text, level=level)

    def paragraph(self, text, size=10, bold=False, italic=False):
        p = self.doc.add_paragraph()
        lines = str(text).split("\n")
        for i, line in enumerate(lines):
            run = p.add_run(line)
            run.font.name = FONT
            run.font.size = self.Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            if i < len(lines) - 1:
                run.add_break()
        return p

    # -- tables --------------------------------------------------------------
    def shade(self, cell, hex_colour):
        shd = self.Oxml("w:shd")
        shd.set(self.qn("w:fill"), hex_colour.lstrip("#"))
        cell._tc.get_or_add_tcPr().append(shd)

    def table(self, header, rows, weights=None, header_fill=None):
        """A styled table: ink header, banded rows. ``header`` may be None
        for a table with no header row (see :meth:`kv_table`)."""
        ncols = len(header) if header else (len(rows[0]) if rows else 0)
        if not ncols:
            return None
        nrows = len(rows) + (1 if header else 0)
        t = self.doc.add_table(rows=nrows, cols=ncols)
        t.style = "Table Grid"
        fill = header_fill or self.ink
        allrows = ([list(header)] if header else []) + [list(r) for r in rows]
        for ri, row in enumerate(allrows):
            is_head = bool(header) and ri == 0
            for ci, val in enumerate(row):
                cell = t.cell(ri, ci)
                cell.text = str(val)
                p = cell.paragraphs[0]
                run = p.runs[0] if p.runs else p.add_run("")
                run.font.name = FONT
                run.font.size = self.Pt(9)
                run.font.bold = is_head
                if is_head:
                    run.font.color.rgb = self._rgb("#FFFFFF")
                self.shade(cell, fill if is_head else
                          (self.alt if ri % 2 == 0 else "#FFFFFF"))
        if weights:
            total = float(sum(weights))
            widths = [self.Inches(PAGE_W_IN * w / total) for w in weights]
            for row in t.rows:
                for cell, w in zip(row.cells, widths):
                    cell.width = w
        return t

    def kv_table(self, pairs):
        """Label/value pairs as a borderless two-column table."""
        t = self.doc.add_table(rows=len(pairs), cols=2)
        for ri, (k, v) in enumerate(pairs):
            c0, c1 = t.cell(ri, 0), t.cell(ri, 1)
            c0.text, c1.text = str(k), str(v)
            for run in c0.paragraphs[0].runs:
                run.font.bold = True
                run.font.name = FONT
                run.font.size = self.Pt(9)
            for run in c1.paragraphs[0].runs:
                run.font.name = FONT
                run.font.size = self.Pt(9)
        t.columns[0].width = self.Inches(1.6)
        return t

    # -- pictures --------------------------------------------------------------
    def picture(self, png_bytes, width_in=PAGE_W_IN):
        self.doc.add_picture(io.BytesIO(png_bytes), width=self.Inches(width_in))

    # -- fields: the TOC and the page-number footer -----------------------------
    def _field(self, paragraph, instr, placeholder=""):
        run = paragraph.add_run()
        r = run._r
        begin = self.Oxml("w:fldChar")
        begin.set(self.qn("w:fldCharType"), "begin")
        instr_el = self.Oxml("w:instrText")
        instr_el.set(self.qn("xml:space"), "preserve")
        instr_el.text = instr
        sep = self.Oxml("w:fldChar")
        sep.set(self.qn("w:fldCharType"), "separate")
        t_el = self.Oxml("w:t")
        t_el.text = placeholder
        end = self.Oxml("w:fldChar")
        end.set(self.qn("w:fldCharType"), "end")
        r.append(begin)
        r.append(instr_el)
        r.append(sep)
        r.append(t_el)
        r.append(end)
        return run

    def toc(self):
        p = self.doc.add_paragraph()
        self._field(p, 'TOC \\o "1-2" \\h \\z \\u',
                    "Right-click here and choose “Update Field” to "
                    "fill in the contents.")

    def footer(self, title):
        section = self.doc.sections[0]
        p = section.footer.paragraphs[0]
        p.alignment = self.ALIGN.CENTER
        run = p.add_run(f"{title}   |   Page ")
        run.font.size, run.font.name = self.Pt(8), FONT
        self._field(p, "PAGE", "1")
        run = p.add_run(" of ")
        run.font.size, run.font.name = self.Pt(8), FONT
        self._field(p, "NUMPAGES", "1")


# -- sections -----------------------------------------------------------------
def _cover_section(d, details, logo, cover, cover_data, notes):
    art = covers.art(cover, "docx", cover_data)
    if art.note and notes is not None:
        notes.append(art.note)
    if art.data:
        d.doc.add_picture(io.BytesIO(art.data), width=d.Inches(art.width))
    if logo and os.path.isfile(logo):
        try:
            d.doc.add_picture(logo, width=d.Inches(2.2))
        except Exception:
            pass                        # a bad logo must not stop the document
    title = (details.get("title") or "").strip() or "Experiment report"
    d.heading(title, level=0)
    date = (details.get("date") or "").strip() \
        or datetime.date.today().isoformat()
    rows = [(label, (details.get(key) or "").strip()) for label, key in
            (("Customer", "customer"), ("Reference", "reference"),
             ("Operator", "operator"))] + [("Date", date)]
    rows = [(k, v) for k, v in rows if v]
    if rows:
        d.kv_table(rows)
    return True


def _contents_section(d, notes):
    d.heading("Contents", level=1)
    d.toc()
    if notes is not None:
        notes.append("The Word document's Contents is a live field: in "
                     "Word, right-click it and choose “Update Field” "
                     "(or select it and press F9) to fill in the page "
                     "numbers.")
    return True


def _glance_section(d, pairs):
    """"At a glance": the same single-line facts as the PDF and the slides."""
    if not pairs:
        return False
    d.heading("At a glance", level=1)
    d.kv_table([(k, v) for k, v in pairs])
    return True


def _timing_section(d, data):
    """"Timing": the rows, the counting time of each sample when that says
    something, and what the files do not record."""
    if not data:
        return False
    d.heading("Timing", level=1)
    d.kv_table([(k, v) for k, v in data.rows])
    if data.by_sample:
        d.paragraph("Counting time by sample", bold=True, size=9)
        d.table(["Sample", "File", "Counting time"],
                [list(r) for r in data.by_sample], [3.0, 3.0, 1.5])
    for note in data.notes:
        d.paragraph(note, size=8, italic=True)
    return True


def _text_section(d, heading, paras):
    if not paras:
        return False
    d.heading(heading, level=1)
    for para in paras:
        d.paragraph(para)
    return True


def _appendix(d):
    """The small "Appendix" label before the first audit section (the PDF
    has the same). A plain paragraph kept with the heading below it, not a
    heading, so it stays out of the Contents field."""
    d.paragraph("Appendix", size=9, bold=True).paragraph_format \
        .keep_with_next = True


def _files_section(d, file_rows, sha="short", appendix=False):
    if not file_rows:
        return False
    if appendix:
        _appendix(d)
    d.heading("Data files", level=1)
    with_sha = sha != "none"
    header = ["File", "Format", "Regions", "Size"] + (["SHA-256"]
                                                       if with_sha else [])
    rows = [[r.get("name", ""), r.get("format", ""), str(r.get("regions", "")),
             _size(r.get("size", 0))]
            + ([(r.get("sha256") or "")[:16]] if with_sha else [])
            for r in file_rows]
    weights = [4.2, 3.0, 0.9, 1.1, 2.3] if with_sha else [6.4, 4.0, 1.0, 1.5]
    d.table(header, rows, weights)
    return True


def _metadata_section(d, docs, skip=(), appendix=False):
    written = False
    for parser in docs:
        if reportspec.doc_key(parser) in skip:
            continue
        samples = parser.samples_metadata()
        if not samples:
            continue
        name = os.path.basename(parser.path or "experiment")
        if not written:
            if appendix:
                _appendix(d)
            d.heading("Acquisition metadata", level=1)
        written = True
        d.heading(name, level=2)
        lay = metasummary.layout_file(samples)
        if lay.common:
            d.paragraph("Common to every region", bold=True, size=9)
            d.kv_table([(metasummary.SHORT.get(k, k), v)
                       for k, v in lay.common])
        for sl in lay.samples:
            label = sl.name + (f"  ({sl.n_regions} regions)"
                               if sl.n_regions else "")
            d.paragraph(label, bold=True, size=9)
            if sl.line:
                d.kv_table([(metasummary.SHORT.get(k, k), v)
                           for k, v in sl.line])
            if sl.rows:
                weights = metasummary.column_weights(sl.columns)
                d.table(sl.columns,
                       [[r.get(c, "") for c in sl.columns] for r in sl.rows],
                       weights)
            for strip_title, entries in sl.strips:
                d.paragraph(strip_title, bold=True, size=8)
                d.table(["Level", "Etch (s)", "Acquired"],
                       [list(e) for e in entries], [0.5, 0.8, 2.2])
    return written


def _results_section(d, results, skip=()):
    samples = results.chosen(skip) if results else []
    if not samples:
        return False
    d.heading("Quantification", level=1)
    d.paragraph(results.method, size=8)
    for s in samples:
        d.heading(s.label, level=2)
        if s.casaxps and not s.levels:
            for title, header, rows, weights in resultspages.casaxps_tables(
                    s.casaxps):
                d.paragraph(title, bold=True, size=9)
                d.table(header, [list(r) for r in rows], weights)
            continue
        if not s.is_profile:
            rows = resultspages.composition_cells(s.levels[0])
            cpng = resultspages.composition_png(s.levels[0], size=(6.5, 3.0),
                                                dpi=200)
            if cpng:
                d.picture(cpng)
            d.table(resultspages.composition_header(s.levels[0]),
                    [c for _k, c in rows],
                   [3.8, 2.8, 1.5, 3.0, 2.4, 1.6, 1.5, 1.4])
            if s.dparam:
                title, header, rows, weights = resultspages.dparam_table(
                    s.dparam)
                d.paragraph(title, bold=True, size=9)
                d.table(header, [list(r) for r in rows], weights)
        else:
            png = resultspages.profile_png(s, size=(6.5, 3.0), dpi=200)
            if png:
                d.picture(png)
            header, rows = resultspages.profile_cells(s)
            d.table(header, rows, [1.0] * len(header))
    for note in results.notes_for(samples):
        d.paragraph("• " + note, size=8)
    return True


def _figures_section(d, figures, render_figure, skip=()):
    chosen = [(i, f) for i, f in enumerate(figures, 1)
             if reportspec.figure_id(f, i) not in skip]
    if not chosen:
        return False
    d.heading("Figures", level=1)
    wrote = False
    for number, fig in chosen:
        images = render_figure(number, fig) or []
        if not images:
            continue
        name = (fig.get("name") or "").strip()
        d.heading(f"Figure {number}" + (f" — {name}" if name else ""),
                 level=2)
        for png in images:
            d.picture(png)
        caption = (fig.get("caption") or "").strip()
        if caption:
            d.paragraph(caption, italic=True, size=9)
        wrote = True
    return wrote


def _images_section(d, image_pages):
    pages = image_pages() if image_pages else []
    if not pages:
        return False
    d.heading("Camera pictures and SnapMaps", level=1)
    for pg in pages:
        d.heading(pg["title"], level=2)
        d.picture(pg["png"])
        if pg.get("notes"):
            d.paragraph(pg["notes"], size=8)
    return True


def build_document(path, details, logo, file_rows, docs, figures,
                   render_figure, image_pages=None, spec=None,
                   cover_data=None, notes=None, results=None, glance=None,
                   timing=None):
    """Write the .docx to ``path``; returns the number of sections written.

    ``spec`` (see ``reportspec``) says which sections go in, in which order,
    which figures, files and pictures, and the cover picture and accent
    colour; without it every section is included in the default order.
    ``cover_data`` is ``(energy, counts)`` for the "your data" cover; a
    problem with the cover picture is appended to ``notes``, along with a
    reminder that the Contents field needs updating in Word.
    ``figures``: ``[{"name", "caption", "state"}]``; ``render_figure(number,
    figure)`` returns one PNG (bytes) per page of that figure.
    ``image_pages()`` returns the camera and SnapMap pictures as
    ``[{"title", "png", "notes"}]`` (None: none). ``results`` is the
    ``resultspages.Results`` the Quantification section is made from.
    ``glance`` is the list of ``(label, value)`` facts of "At a glance" and
    ``timing`` the ``glance.Timing`` of "Timing" (None: left out)."""
    if spec is None:
        spec = reportspec.default_spec()
    items = reportspec.active(spec)
    ids = [sid for sid, _skip in items]
    title = (details.get("title") or "").strip() or "Experiment report"
    cover = reportspec.cover_of(spec)
    d = _Doc(cover["accent"])
    methods = (details.get("methods") or "").strip()
    cal = (details.get("calibration") or "").strip()
    # the calibration statement rides on the summary when both are in (as the
    # PDF and the deck do), else it gets a heading of its own; never twice
    cal_text = cal if cal and not ("methods" in ids and cal in methods) else ""
    written = 0
    first_audit = next((i for i in ids if i in ("metadata", "files")), None)
    for sid, skip in items:
        if sid == "cover":
            written += _cover_section(d, details, logo, cover, cover_data,
                                      notes)
        elif sid == "contents":
            written += _contents_section(d, notes)
        elif sid == "glance":
            written += _glance_section(d, glance)
        elif sid == "summary":
            paras = _paragraphs(details.get("summary"))
            if cal_text and "calibration" in ids:
                paras = paras + ["Energy calibration: " + cal_text]
            written += _text_section(d, "Summary", paras)
        elif sid == "results":
            written += _results_section(d, results, skip)
        elif sid == "methods":
            written += _text_section(d, "Methods", _paragraphs(methods))
        elif sid == "timing":
            written += _timing_section(d, timing)
        elif sid == "calibration":
            if cal_text and "summary" not in ids:
                written += _text_section(d, "Energy calibration", [cal_text])
        elif sid == "files":
            written += _files_section(d, file_rows,
                                      reportspec.option(spec, "sha"),
                                      appendix=sid == first_audit)
        elif sid == "metadata":
            written += _metadata_section(d, docs, skip,
                                         appendix=sid == first_audit)
        elif sid == "images" and image_pages is not None:
            written += _images_section(d, image_pages)
        elif sid == "figures" and figures:
            written += _figures_section(d, figures, render_figure, skip)
    if not written:
        raise DocxError("Nothing to put in the document: choose at least "
                        "one section that has content.")
    d.footer(title)
    core = d.doc.core_properties
    core.title = title
    core.author = (details.get("operator") or "").strip() or appinfo.NAME
    d.doc.save(path)
    return written
