"""PowerPoint export of an experiment: a 16:9 deck built from the same data as
the PDF report.

Slides: title (logo, details), contents, summary, quantification (composition
tables, depth-profile charts; see ``resultspages``), data files, metadata (the
compact layout of ``metasummary.layout_file`` as native tables), the camera
pictures and SnapMaps (one picture per slide, see ``imagepages``), then one
slide per figure page (a picture, an editable caption, speaker notes describing
the look), in the order of the ``reportspec`` spec.

The slides are built in that order, each noting its section; a last pass then
adds the **divider** slides (before a section of ``DIVIDER_MIN`` slides or more,
unless the spec says never) and the **contents** slide(s) (sections and their
slide numbers, where the spec puts them), moves them into place, and stamps
every slide but the title and the dividers with "title | section | n of N", so
the numbers on the contents slide are the real ones and each entry jumps to
its slide on click. The accent colour is the cover's. Needs python-pptx; no
Tk here.
"""

from __future__ import annotations

import datetime
import io
import os

import appinfo
import covers
import metasummary
import panelview
import reportspec
import resultspages

SECTIONS = ("title", "files", "metadata", "images", "figures")

SLIDE_W, SLIDE_H = 13.333, 7.5          # inches (16:9)
MARGIN = 0.6
BODY_W = SLIDE_W - 2 * MARGIN
FIGURE_SIZE = (12.1, 4.95)              # inches, for the picture on a figure slide
GREY = (0x6B, 0x77, 0x85)
INK = (0x1A, 0x21, 0x27)
FONT = "Calibri"

DIVIDER_MIN = 5                         # slides in a section that earn a divider
CONTENTS_ROWS = 16                      # lines on a contents slide
RESULT_ROWS = 15                        # table rows on a quantification slide
ROW_H = 0.29                            # table row height (10 pt text)
TABLE_TOP = 1.25
BOTTOM = SLIDE_H - 0.55                 # keep clear of the footer


class PptxError(Exception):
    """The deck could not be built (message is user-facing)."""


def _modules():
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.action import ActionSetting
    except ImportError:
        raise PptxError("PowerPoint export needs python-pptx "
                        "(pip install python-pptx).")
    return (Presentation, Inches, Pt, RGBColor, MSO_ANCHOR, PP_ALIGN,
            MSO_SHAPE, ActionSetting)


# -- pure helpers ----------------------------------------------------------------
def paragraphs(text):
    """Blank-line separated paragraphs of free text (single newlines kept)."""
    return [b.strip("\n") for b in (text or "").replace("\r\n", "\n")
            .split("\n\n") if b.strip()]


def chunk_paragraphs(paras, max_lines=11, chars_per_line=92):
    """Split paragraphs over slides so each holds about ``max_lines`` lines
    (long paragraphs are cut at word boundaries)."""
    def lines(p):
        return sum(max(1, -(-len(part) // chars_per_line))
                   for part in p.split("\n")) + 1     # + spacing after

    pieces = []
    for p in paras:
        if lines(p) <= max_lines:
            pieces.append(p)
            continue
        cur = ""
        cap = (max_lines - 1) * chars_per_line
        words = []
        for word in p.split(" "):          # an unbreakable run is cut too
            while len(word) > cap:
                words.append(word[:cap])
                word = word[cap:]
            words.append(word)
        for word in words:
            if lines(cur + " " + word) > max_lines and cur:
                pieces.append(cur)
                cur = word
            else:
                cur = (cur + " " + word) if cur else word
        if cur:
            pieces.append(cur)
    slides, cur, used = [], [], 0
    for p in pieces:
        n = lines(p)
        if cur and used + n > max_lines:
            slides.append(cur)
            cur, used = [], 0
        cur.append(p)
        used += n
    if cur:
        slides.append(cur)
    return slides


def look_notes(state):
    """Speaker-notes text describing a saved look."""
    st = state or {}
    parts = [f"View: {st.get('view_mode', 'Stack')}",
             f"grouped by {st.get('group_by', 'element name').lower()}",
             "energy ratio axis (E/E0, ion scattering)"
             if st.get("energy_scale") == "Energy ratio"
             else f"{st.get('energy_scale', 'Binding').lower()} energy axis"]
    if st.get("norm") not in (None, "None"):
        parts.append(f"normalised: {st['norm']}")
    if st.get("colour_scale") not in (None, "Theme default"):
        parts.append(f"colour scale: {st['colour_scale']}"
                     + (" (reversed)" if st.get("colour_reverse") else ""))
    if st.get("view_mode") in ("Waterfall 3D", "Heatmap"):
        parts.append(f"z axis: {st.get('z_axis', 'Auto')}")
    ticked = st.get("ticked") or []
    names = metasummary.compact_labels(
        [f"{t.get('name', '')}" for t in ticked], limit=12)
    text = "; ".join(parts) + "."
    defaults = {"view": st.get("view_mode", "Stack"),
                "norm": st.get("norm", "None"),
                "offset": st.get("offset", 0.6),
                "z_axis": st.get("z_axis", "Auto"),
                "reverse": bool(st.get("reverse")),
                "fit_show": st.get("fit_show") or {}}
    own = [f"{label}: {panelview.describe(ov, defaults)}"
           for label, ov in panelview.sanitise_all(
               st.get("panel_views")).items()]
    if own:
        text += "\nPanels with their own view: " + "; ".join(own) + "."
    if ticked:
        text += f"\nSpectra shown ({len(ticked)}): {names}."
    return text


def pack_items(items, budget):
    """Greedy page packing: ``items`` is ``[(height, payload)]``; returns
    ``[[payload, ...], ...]`` with each page's total height <= ``budget``
    (an oversize item gets a page to itself)."""
    pages, cur, used = [], [], 0.0
    for h, payload in items:
        if cur and used + h > budget:
            pages.append(cur)
            cur, used = [], 0.0
        cur.append(payload)
        used += h
    if cur:
        pages.append(cur)
    return pages


# -- drawing -----------------------------------------------------------------------
class _Deck:
    def __init__(self, title, accent=""):
        (self.Presentation, self.Inches, self.Pt, self.RGB, self.ANCHOR,
         self.ALIGN, self.SHAPE, self.Action) = _modules()
        self.prs = self.Presentation()
        self.prs.slide_width = self.Inches(SLIDE_W)
        self.prs.slide_height = self.Inches(SLIDE_H)
        self.title = title
        # the cover's accent: the bar and rules as it is, text and table
        # headers darkened (so a pale accent still reads on white)
        self.accent = _rgb(covers.valid_accent(accent) or covers.DEFAULT_ACCENT)
        self.ink = _rgb(covers.mix(_hex(self.accent), "#000000", 0.2))
        self.alt = _rgb(covers.tint(_hex(self.accent), 0.94))
        self.section = "cover"          # what the slides being added belong to
        self.slides = []                # (slide, {"sid", "child", "kind"})
        self.contents_at = None         # index of the slide it goes before

    # low level
    def rgb(self, t):
        return self.RGB(*t)

    def bar(self, slide):
        shp = slide.shapes.add_shape(self.SHAPE.RECTANGLE, 0, 0,
                                     self.Inches(SLIDE_W), self.Inches(0.14))
        shp.fill.solid()
        shp.fill.fore_color.rgb = self.rgb(self.accent)
        shp.line.fill.background()

    def text(self, slide, x, y, w, h, paras, size=14, bold=False,
             color=INK, align=None, space_after=6):
        box = slide.shapes.add_textbox(self.Inches(x), self.Inches(y),
                                       self.Inches(w), self.Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = self.Pt(space_after)
            if align:
                p.alignment = align
            r = p.add_run()
            r.text = para
            r.font.size = self.Pt(size)
            r.font.bold = bold
            r.font.name = FONT
            r.font.color.rgb = self.rgb(color)
        return box

    def note(self, slide, sid, child=None, kind="content"):
        """Record what a slide is (section, sub-entry) for the last pass."""
        self.slides.append((slide, {"sid": sid, "child": child, "kind": kind}))

    def titled_slide(self, title):
        """A 'Title Only' slide with the title restyled to fit 16:9."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[5])
        self.bar(slide)
        t = slide.shapes.title
        t.left, t.top = self.Inches(MARGIN), self.Inches(0.35)
        t.width, t.height = self.Inches(BODY_W), self.Inches(0.75)
        tf = t.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = self.ANCHOR.MIDDLE
        tf.text = title
        p = tf.paragraphs[0]
        p.alignment = self.ALIGN.LEFT
        for r in p.runs:
            r.font.size = self.Pt(26 if len(title) <= 52 else 20)
            r.font.bold = True
            r.font.name = FONT
            r.font.color.rgb = self.rgb(self.ink)
        return slide

    def content_slide(self, title, child=None):
        """A titled slide of the section being built; ``child`` names the
        file or figure it starts (an entry under the section in the contents)."""
        slide = self.titled_slide(title)
        self.note(slide, self.section, child)
        return slide

    def table(self, slide, x, y, w, weights, header, rows, size=10,
              row_h=ROW_H, header_fill=None):
        header_fill = header_fill or self.ink
        nrows = len(rows) + (1 if header else 0)
        shape = slide.shapes.add_table(nrows, len(weights), self.Inches(x),
                                       self.Inches(y), self.Inches(w),
                                       self.Inches(row_h * nrows))
        tbl = shape.table
        tbl.horz_banding = False
        total = float(sum(weights))
        widths = [self.Inches(w * v / total) for v in weights]
        widths[-1] = self.Inches(w) - sum(widths[:-1])     # exact total
        for col, wd in zip(tbl.columns, widths):
            col.width = wd
        for r in tbl.rows:
            r.height = self.Inches(row_h)
        allrows = ([list(header)] if header else []) + rows
        for ri, row in enumerate(allrows):
            is_head = bool(header) and ri == 0
            for ci, val in enumerate(row):
                cell = tbl.cell(ri, ci)
                cell.margin_left = cell.margin_right = self.Inches(0.05)
                cell.margin_top = cell.margin_bottom = self.Inches(0.02)
                cell.vertical_anchor = self.ANCHOR.MIDDLE
                cell.fill.solid()
                cell.fill.fore_color.rgb = self.rgb(
                    header_fill if is_head else
                    (self.alt if (ri % 2 == 0) else (255, 255, 255)))
                tf = cell.text_frame
                tf.word_wrap = True
                tf.text = str(val)
                for p in tf.paragraphs:
                    for r in p.runs:
                        r.font.size = self.Pt(size)
                        r.font.name = FONT
                        r.font.bold = is_head
                        r.font.color.rgb = self.rgb((255, 255, 255)
                                                    if is_head else INK)
        return shape


def _rgb(hex_colour):
    h = hex_colour.lstrip("#")
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


def _hex(rgb):
    return "#%02X%02X%02X" % tuple(rgb)


# -- slide builders ------------------------------------------------------------------
def _title_slide(deck, details, logo, art=None):
    slide = deck.prs.slides.add_slide(deck.prs.slide_layouts[0])
    deck.bar(slide)
    if art is not None and art.data:           # a strip along the bottom
        slide.shapes.add_picture(
            io.BytesIO(art.data), 0, deck.Inches(SLIDE_H - art.height),
            width=deck.Inches(SLIDE_W))
    ph = list(slide.placeholders)
    title = (details.get("title") or "").strip() or "Experiment report"
    t = slide.shapes.title
    t.left, t.top = deck.Inches(0.8), deck.Inches(2.3)
    t.width, t.height = deck.Inches(SLIDE_W - 1.6), deck.Inches(1.5)
    t.text_frame.word_wrap = True
    t.text_frame.vertical_anchor = deck.ANCHOR.BOTTOM
    t.text_frame.text = title
    for p in t.text_frame.paragraphs:
        p.alignment = deck.ALIGN.LEFT
        for r in p.runs:
            r.font.size = deck.Pt(38)
            r.font.bold = True
            r.font.name = FONT
            r.font.color.rgb = deck.rgb(deck.ink)
    date = (details.get("date") or "").strip() \
        or datetime.date.today().isoformat()
    lines = [f"{label}: {v}" for label, v in
             (("Customer", details.get("customer")),
              ("Reference", details.get("reference")),
              ("Operator", details.get("operator")), ("Date", date))
             if v and str(v).strip()]
    sub = [p for p in ph if p.placeholder_format.idx == 1]
    if sub:
        s = sub[0]
        s.left, s.top = deck.Inches(0.8), deck.Inches(4.0)
        s.width, s.height = deck.Inches(SLIDE_W - 1.6), deck.Inches(1.8)
        tf = s.text_frame
        tf.word_wrap = True
        tf.text = lines[0] if lines else ""
        for line in lines[1:]:
            tf.add_paragraph().text = line
        for p in tf.paragraphs:
            p.alignment = deck.ALIGN.LEFT
            p.space_after = deck.Pt(4)
            for r in p.runs:
                r.font.size = deck.Pt(20)
                r.font.name = FONT
                r.font.color.rgb = deck.rgb(GREY)
    if logo and os.path.isfile(logo):
        try:
            pic = slide.shapes.add_picture(logo, deck.Inches(0.8),
                                           deck.Inches(0.7),
                                           width=deck.Inches(3.2))
            if pic.height > deck.Inches(1.3):          # keep it compact
                ratio = deck.Inches(1.3) / pic.height
                pic.height = deck.Inches(1.3)
                pic.width = int(pic.width * ratio)
        except Exception:
            pass                                   # a bad logo must not stop the deck
    return slide


def _text_slides(deck, title, paras, size=16):
    """Slides of a heading and free text (paginated); ``paras`` is a list of
    paragraphs."""
    for i, page in enumerate(chunk_paragraphs(paras)):
        slide = deck.content_slide(title + (" (continued)" if i else ""))
        deck.text(slide, MARGIN, 1.3, BODY_W, BOTTOM - 1.3, page,
                  size=size, space_after=10)


def _glance_slides(deck, pairs):
    """One slide of single-line facts (the same lines as the PDF)."""
    rows = [[str(k), str(v)] for k, v in pairs]
    per = int((BOTTOM - TABLE_TOP) / ROW_H)
    for i in range(0, max(1, len(rows)), per):
        slide = deck.content_slide(
            "At a glance" + (" (continued)" if i else ""))
        _kv_table(deck, slide, TABLE_TOP, rows[i:i + per], 2)


def _timing_slides(deck, data):
    """The Timing rows (and what the files do not record) on one slide, then
    the counting time of each sample as a table over as many slides as it
    needs."""
    rows = [[str(k), str(v)] for k, v in data.rows]
    slide = deck.content_slide("Timing")
    _kv_table(deck, slide, TABLE_TOP, rows, 2)
    if data.notes:
        top = TABLE_TOP + ROW_H * len(rows) + 0.3
        deck.text(slide, MARGIN, top, BODY_W, max(0.5, BOTTOM - top),
                  list(data.notes), size=12, space_after=6)
    per = int((BOTTOM - TABLE_TOP) / ROW_H) - 1
    for i in range(0, len(data.by_sample), per):
        slide = deck.content_slide(
            "Timing: counting time by sample"
            + (" (continued)" if i else ""))
        deck.table(slide, MARGIN, TABLE_TOP, BODY_W, [3.0, 3.0, 1.5],
                   ["Sample", "File", "Counting time"],
                   [list(r) for r in data.by_sample[i:i + per]])


def _files_slides(deck, file_rows, sha="short"):
    per = int((BOTTOM - TABLE_TOP) / ROW_H) - 1
    with_sha = sha != "none"
    for i in range(0, max(1, len(file_rows)), per):
        chunk = file_rows[i:i + per]
        slide = deck.content_slide(
            "Data files" + (" (continued)" if i else ""))
        rows = [[r.get("name", ""), r.get("format", ""),
                 str(r.get("regions", "")), _size(r.get("size", 0))]
                + ([(r.get("sha256") or "")[:16]] if with_sha else [])
                for r in chunk]
        deck.table(slide, MARGIN, TABLE_TOP, BODY_W,
                   [4.2, 3.0, 0.9, 1.1, 2.3] if with_sha
                   else [6.4, 4.0, 1.0, 1.5],
                   ["File", "Format", "Regions", "Size"]
                   + (["SHA-256"] if with_sha else []), rows)


def _size(n):
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024.0


def _metadata_slides(deck, docs, skip=()):
    budget = BOTTOM - TABLE_TOP
    for parser in docs:
        if reportspec.doc_key(parser) in skip:
            continue
        samples = parser.samples_metadata()
        if not samples:
            continue
        name = os.path.basename(parser.path or "experiment")
        lay = metasummary.layout_file(samples)
        items = []                                # (height, draw-callable)
        strips = []

        if lay.common:
            kv = [(metasummary.SHORT.get(k, k), v) for k, v in lay.common]
            rows = [sum(([k, v] for k, v in kv[i:i + 3]), [])
                    for i in range(0, len(kv), 3)]
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            h = 0.32 + ROW_H * len(rows)
            items.append((h, ("common", rows, width)))

        for sl in lay.samples:
            sub = [(metasummary.SHORT.get(k, k), v) for k, v in sl.line]
            line_rows = [sum(([k, v] for k, v in sub[i:i + 3]), [])
                         for i in range(0, len(sub), 3)]
            if line_rows:
                w_ = max(len(r) for r in line_rows)
                line_rows = [r + [""] * (w_ - len(r)) for r in line_rows]
            head_h = 0.36 + ROW_H * len(line_rows)
            max_rows = max(3, int((budget - head_h) / ROW_H) - 1)
            for i in range(0, max(1, len(sl.rows)), max_rows):
                chunk = sl.rows[i:i + max_rows]
                first = i == 0
                h = (head_h if first else 0.36) + ROW_H * (len(chunk) + 1)
                items.append((h, ("sample", sl, chunk, first, line_rows)))
            for title, entries in sl.strips:
                strips.append((f"{sl.name}: {title}", entries))

        for pi, page in enumerate(pack_items(items, budget)):
            slide = deck.content_slide(
                f"Acquisition metadata – {name}"
                + (" (continued)" if pi else ""), None if pi else name)
            y = TABLE_TOP
            for item in page:
                if item[0] == "common":
                    _, rows, width = item
                    deck.text(slide, MARGIN, y, BODY_W, 0.3,
                              ["Common to every region"], size=12,
                              bold=True, color=deck.ink, space_after=0)
                    y += 0.32
                    _kv_table(deck, slide, y, rows, width)
                    y += ROW_H * len(rows)
                else:
                    _, sl, chunk, first, line_rows = item
                    label = sl.name + (
                        f"  ({sl.n_regions} regions)" if first else
                        "  (continued)")
                    deck.text(slide, MARGIN, y, BODY_W, 0.3, [label],
                              size=12, bold=True, color=deck.ink, space_after=0)
                    y += 0.36
                    if first and line_rows:
                        _kv_table(deck, slide, y, line_rows,
                                  max(len(r) for r in line_rows))
                        y += ROW_H * len(line_rows)
                    weights = metasummary.column_weights(sl.columns)
                    deck.table(slide, MARGIN, y, BODY_W, weights, sl.columns,
                               [[r.get(c, "") for c in sl.columns]
                                for r in chunk])
                    y += ROW_H * (len(chunk) + 1)

        for title, entries in strips:                # per-level details
            per_slide = 60
            for i in range(0, len(entries), per_slide):
                part = entries[i:i + per_slide]
                slide = deck.content_slide(
                    f"{title} – {name}" + (" (continued)" if i else ""))
                percol = -(-len(part) // 3)
                grid = []
                for r in range(percol):
                    row = []
                    for c in range(3):
                        idx = c * percol + r
                        row += list(part[idx]) if idx < len(part) \
                            else ["", "", ""]
                    grid.append(row)
                deck.table(slide, MARGIN, TABLE_TOP, BODY_W,
                           [0.5, 0.8, 2.2] * 3,
                           ["Level", "Etch (s)", "Acquired"] * 3, grid,
                           size=9, row_h=0.26)


def _kv_table(deck, slide, y, rows, width):
    """Key/value pairs laid out across (label, value, label, value ...)."""
    pairs = width // 2
    weights = [1.0, 1.6] * pairs
    shape = deck.table(slide, MARGIN, y, BODY_W, weights, None, rows)
    tbl = shape.table
    for ri, row in enumerate(rows):
        for ci in range(0, width, 2):
            cell = tbl.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = deck.rgb((255, 255, 255))
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.bold = True
                    r.font.color.rgb = deck.rgb(deck.ink)
            for cj in (ci + 1,):
                c2 = tbl.cell(ri, cj)
                c2.fill.solid()
                c2.fill.fore_color.rgb = deck.rgb((255, 255, 255))
    return shape


def _figure_slides(deck, figures, render_images):
    for number, fig in enumerate(figures, 1):
        images = render_images(number, fig)
        caption = (fig.get("caption") or "").strip()
        for pi, png in enumerate(images):
            title = f"Figure {number} – {fig.get('name', '')}"
            if pi:
                title += " (continued)"
            slide = deck.content_slide(title, None if pi else title)
            _fit_picture(deck, slide, png)
            if caption and pi == 0:
                deck.text(slide, MARGIN, 6.15, BODY_W, 0.9,
                          paragraphs(caption)[:3], size=13, space_after=3)
            notes = slide.notes_slide.notes_text_frame
            notes.text = ((caption + "\n\n") if caption else "") \
                + look_notes(fig.get("state"))


def _image_slides(deck, pages):
    """One slide per camera sheet / SnapMap site: ``pages`` is
    ``[{"title", "png", "notes"}]`` (pictures sized ``FIGURE_SIZE``). Each
    page is its own entry under Pictures in the contents, the page's own
    title being unique enough already (a sheet's "(n of M)", a site's own
    name)."""
    for pg in pages:
        slide = deck.content_slide(pg["title"], pg["title"])
        _fit_picture(deck, slide, pg["png"])
        slide.notes_slide.notes_text_frame.text = pg.get("notes", "")


def _fit_picture(deck, slide, png):
    """Place a picture sized ``FIGURE_SIZE`` under the title, centred, and
    shrink it if it came out taller."""
    width, height = FIGURE_SIZE
    pic = slide.shapes.add_picture(
        io.BytesIO(png), deck.Inches((SLIDE_W - width) / 2),
        deck.Inches(1.15), width=deck.Inches(width))
    if pic.height > deck.Inches(height):
        ratio = deck.Inches(height) / pic.height
        pic.height = deck.Inches(height)
        pic.width = int(pic.width * ratio)
        pic.left = int((deck.Inches(SLIDE_W) - pic.width) / 2)
    return pic


def _fmt(v, spec=".2f"):
    return "" if v is None else format(v, spec)


def _right_align(deck, shape, first):
    for row in shape.table.rows:
        for ci in range(first, len(shape.table.columns)):
            for p in row.cells[ci].text_frame.paragraphs:
                p.alignment = deck.ALIGN.RIGHT


def _table_slides(deck, s, finish, title, header, rows, weights):
    """One or more table slides (``RESULT_ROWS`` a slide) for ``rows``."""
    for i in range(0, len(rows), RESULT_ROWS):
        chunk = rows[i:i + RESULT_ROWS]
        slide = deck.content_slide(
            f"Quantification – {s.label}: {title}"
            + (" (continued)" if i else ""),
            None if i else s.label)
        shape = deck.table(slide, MARGIN, TABLE_TOP, BODY_W, weights,
                           header, chunk, size=11, row_h=0.32)
        _right_align(deck, shape, len(header) - 1)
        finish(slide)


def _casaxps_slides(deck, s):
    """One or more table slides per CasaXPS-exported table of a sample
    (``s.casaxps``, see ``casaquant.py``): the numbers are CasaXPS's own,
    not recomputed, so the footnote says that instead of a formula."""
    cq = s.casaxps

    def finish(slide):
        deck.text(slide, MARGIN, 6.55, BODY_W, 0.55,
                  [resultspages.CASAXPS_NOTE], size=10, color=GREY,
                  space_after=0)
        slide.notes_slide.notes_text_frame.text = resultspages.CASAXPS_NOTE

    def table_slides(title, header, rows, weights):
        for i in range(0, len(rows), RESULT_ROWS):
            chunk = rows[i:i + RESULT_ROWS]
            slide = deck.content_slide(
                f"Quantification – {s.label}: {title}"
                + (" (continued)" if i else ""),
                None if i else s.label)
            shape = deck.table(slide, MARGIN, TABLE_TOP, BODY_W, weights,
                               header, chunk, size=11, row_h=0.32)
            _right_align(deck, shape, len(header) - 1)
            finish(slide)

    if cq.survey:
        table_slides("survey", ("Element", "%Conc"),
                     [(r["element"], _fmt(r["pct"])) for r in cq.survey],
                     [2.0, 1.0])
    if cq.regions:
        table_slides("regions", ("Name", "Position (eV)", "%At Conc"),
                     [(r["name"], _fmt(r["position"], "g"),
                       _fmt(r["at_pct"])) for r in cq.regions],
                     [2.4, 1.4, 1.2])
    if cq.dparam:
        table_slides("D parameter", ("Name", "FWHM (eV)"),
                     [(r["name"], _fmt(r["fwhm"], "g")) for r in cq.dparam],
                     [2.0, 1.0])


def _results_slides(deck, results, skip=()):
    """The Quantification slides: for each sample its composition (a chart,
    then a table, one depth level) or its depth profile (a chart, then at %
    by level), a footnote on how the numbers are made, and the method and
    notes as speaker notes."""
    samples = results.chosen(skip) if results else []
    if not samples:
        return
    foot = "Atomic % = area / RSF (CasaXPS), per sample and depth level; " \
        "no transmission correction."
    notes = results.notes_for(samples)
    said = results.method + ("\n\n" + "\n".join(notes) if notes else "")

    cur = [foot]                   # the footnote of the sample being laid out

    def finish(slide, extra=None):
        lines = [cur[0]] if extra is None else [cur[0], extra]
        deck.text(slide, MARGIN, 6.55, BODY_W, 0.35 if extra is None else 0.55,
                  lines, size=10, color=GREY, space_after=0)
        slide.notes_slide.notes_text_frame.text = said

    def right(shape, first):
        _right_align(deck, shape, first)

    for s in samples:
        if s.casaxps:
            _casaxps_slides(deck, s)
            continue
        cur[0] = resultspages.CASA_FOOT if s.numbers == "casaxps" else foot
        if not s.is_profile:
            header, rows, weights = resultspages.composition_table(
                s.levels[0], [3.8, 2.8, 1.5, 3.0, 2.4, 1.6, 1.5, 1.4])
            cpng = resultspages.composition_png(s.levels[0], size=FIGURE_SIZE,
                                                dpi=150)
            if cpng:
                slide = deck.content_slide(f"Quantification \u2013 {s.label}",
                                           s.label)
                _fit_picture(deck, slide, cpng)
                finish(slide)
            for i in range(0, len(rows), RESULT_ROWS):
                chunk = rows[i:i + RESULT_ROWS]
                suffix = ": composition table" if cpng else ""
                slide = deck.content_slide(
                    f"Quantification \u2013 {s.label}{suffix}"
                    + (" (continued)" if i else ""),
                    None if (cpng or i) else s.label)
                shape = deck.table(
                    slide, MARGIN, TABLE_TOP, BODY_W, weights, header,
                    [c for _k, c in chunk], size=11, row_h=0.32)
                right(shape, 2)
                for ri, (kind, _c) in enumerate(chunk, 1):
                    if kind == "state":
                        for cell in shape.table.rows[ri].cells:
                            for p in cell.text_frame.paragraphs:
                                for r in p.runs:
                                    r.font.size = deck.Pt(10)
                                    r.font.color.rgb = deck.rgb(GREY)
                finish(slide)
            if s.dparam:
                _table_slides(deck, s, finish, *resultspages.dparam_table(
                    s.dparam))
            continue
        slide = deck.content_slide(
            f"Quantification \u2013 {s.label}: depth profile", s.label)
        png = resultspages.profile_png(s, size=FIGURE_SIZE, dpi=150)
        if png:
            _fit_picture(deck, slide, png)
        finish(slide)
        header, rows = resultspages.profile_cells(s)
        for i in range(0, len(rows), RESULT_ROWS):
            slide = deck.content_slide(
                f"Quantification \u2013 {s.label}: at % by level"
                + (" (continued)" if i else ""))
            shape = deck.table(slide, MARGIN, TABLE_TOP, BODY_W,
                               [1.0] * len(header), header,
                               rows[i:i + RESULT_ROWS], size=11, row_h=0.32)
            right(shape, 1)
            finish(slide)


# -- the last pass: dividers, contents, order, footers -----------------------------
def _entries(deck):
    """``[(section, level, title, slide index)]`` for the contents: one line
    per section, and under it its files or figures when there are two or more.
    The index is into ``deck.slides`` (creation order)."""
    heads, kids, order = {}, {}, []
    for i, (_slide, m) in enumerate(deck.slides):
        if m["kind"] != "content":
            continue
        if m["sid"] not in heads:
            heads[m["sid"]] = i
            order.append(m["sid"])
        if m["child"]:
            kids.setdefault(m["sid"], []).append((m["child"], i))
    out = []
    for sid in order:
        out.append((sid, 1, reportspec.LABELS[sid], heads[sid]))
        if len(kids.get(sid, ())) >= 2:
            out += [(sid, 2, title, i) for title, i in kids[sid]]
    return out


def _divider_slide(deck, sid, count):
    """A full-bleed slide naming a section (and how many slides it has)."""
    slide = deck.prs.slides.add_slide(deck.prs.slide_layouts[6])
    back = slide.shapes.add_shape(deck.SHAPE.RECTANGLE, 0, 0,
                                  deck.Inches(SLIDE_W), deck.Inches(SLIDE_H))
    back.fill.solid()
    back.fill.fore_color.rgb = deck.rgb(deck.ink)
    back.line.fill.background()
    stripe = slide.shapes.add_shape(deck.SHAPE.RECTANGLE, deck.Inches(0.8),
                                    deck.Inches(3.75), deck.Inches(1.4),
                                    deck.Inches(0.07))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = deck.rgb(_rgb(covers.tint(_hex(deck.accent),
                                                           0.55)))
    stripe.line.fill.background()
    deck.text(slide, 0.8, 2.55, SLIDE_W - 1.6, 1.1, [reportspec.LABELS[sid]],
              size=44, bold=True, color=(255, 255, 255))
    deck.text(slide, 0.8, 4.0, SLIDE_W - 1.6, 0.6,
              [f"{reportspec.HINTS[sid]}  •  {count} slides"], size=18,
              color=_rgb(covers.tint(_hex(deck.accent), 0.8)))
    return slide


def _contents_slide(deck, rows, first):
    """One slide of the contents: ``rows`` is ``[(level, title, number,
    target_slide)]``; the title cell jumps to ``target_slide`` on click."""
    slide = deck.titled_slide("Contents" + ("" if first else " (continued)"))
    cells = [[("    " if lvl == 2 else "") + title, str(num)]
             for lvl, title, num, _target in rows]
    shape = deck.table(slide, MARGIN, TABLE_TOP, BODY_W, [11.0, 1.0], None,
                       cells, size=12, row_h=0.32)
    for ri, (lvl, _title, _num, target) in enumerate(rows):
        for ci in (0, 1):
            cell = shape.table.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = deck.rgb(deck.alt if lvl == 1
                                                else (255, 255, 255))
            for p in cell.text_frame.paragraphs:
                if ci:
                    p.alignment = deck.ALIGN.RIGHT
                for r in p.runs:
                    r.font.bold = lvl == 1
                    r.font.size = deck.Pt(12 if lvl == 1 else 11)
                    r.font.color.rgb = deck.rgb(deck.ink if lvl == 1 else INK)
                    if ci == 0 and target is not None:
                        deck.Action(r.font._rPr, r).target_slide = target
    return slide


def _assemble(deck, dividers="auto"):
    """Add the divider and contents slides, put every slide in its place and
    stamp the footers. The slide numbers on the contents are the final ones,
    and each entry jumps to its slide on click."""
    main = len(deck.slides)
    entries = _entries(deck)
    counts = {}
    for _slide, m in deck.slides:
        if m["kind"] == "content":
            counts[m["sid"]] = counts.get(m["sid"], 0) + 1
    firsts = {sid: i for sid, lvl, _t, i in entries if lvl == 1}
    divided = {firsts[sid]: sid for sid in firsts
               if dividers == "auto" and counts[sid] >= DIVIDER_MIN}
    pages = -(-len(entries) // CONTENTS_ROWS) \
        if entries and deck.contents_at is not None else 0

    layout = []            # ("slide", i) | ("divider", sid) | ("contents", j)
    for i in range(main + 1):
        if i == deck.contents_at:
            layout += [("contents", j) for j in range(pages)]
        if i in divided:
            layout.append(("divider", divided[i]))
        if i < main:
            layout.append(("slide", i))
    number = {item: k for k, item in enumerate(layout, 1)}

    made = []              # the extra slides, in creation order
    for sid in divided.values():
        made.append((("divider", sid),
                     _divider_slide(deck, sid, counts[sid])))

    # every slide a contents entry could jump to (an ordinary slide, or a
    # divided section's own divider -- never a contents slide itself, since
    # one is never among ``entries``); built now, before the contents
    # slide(s) exist, so their own table can carry the links
    target_of = {("slide", i): slide
                for i, (slide, _m) in enumerate(deck.slides)}
    target_of.update(dict(made))

    rows = []
    for sid, lvl, title, i in entries:
        own = ("divider", sid)
        item = own if (lvl == 1 and own in number) else ("slide", i)
        rows.append((lvl, title, number[item], target_of[item]))
    for j in range(pages):
        chunk = rows[j * CONTENTS_ROWS:(j + 1) * CONTENTS_ROWS]
        made.append((("contents", j), _contents_slide(deck, chunk, j == 0)))

    lst = deck.prs.slides._sldIdLst
    elems = list(lst)      # creation order: the slides, then ``made``
    by_item = {("slide", i): elems[i] for i in range(main)}
    by_item.update({item: elems[main + k]
                    for k, (item, _s) in enumerate(made)})
    for el in elems:
        lst.remove(el)
    for item in layout:
        lst.append(by_item[item])

    slide_of = {("slide", i): slide
                for i, (slide, _m) in enumerate(deck.slides)}
    slide_of.update(dict(made))
    total = len(layout)
    for item in layout:
        if item[0] == "divider":
            continue
        label = "Contents"
        if item[0] == "slide":
            meta = deck.slides[item[1]][1]
            if meta["kind"] == "cover":
                continue
            label = reportspec.LABELS.get(meta["sid"], "")
        deck.text(slide_of[item], MARGIN, SLIDE_H - 0.42, BODY_W, 0.3,
                  ["   |   ".join(x for x in (deck.title, label,
                                               f"{number[item]} of {total}")
                                  if x)], size=9, color=GREY)


def build_deck(path, details, logo, file_rows, docs, figures, render_images,
               sections=SECTIONS, image_pages=None, spec=None,
               cover_data=None, notes=None, results=None, glance=None,
               timing=None):
    """Write the .pptx to ``path``; returns the number of slides.

    ``spec`` (see ``reportspec``) says which sections go in, in which order,
    which figures and files, the cover picture and accent colour, and whether
    long sections get a divider slide; without it ``sections`` (the old names)
    do. ``cover_data`` is ``(energy, counts)`` for the "your data" cover; a
    problem with the cover picture is appended to ``notes``.
    ``figures``: ``[{"name", "caption", "state"}]``;
    ``render_images(number, figure)`` returns one PNG (bytes) per page of that
    figure, sized ``FIGURE_SIZE`` inches. ``image_pages()`` returns the camera
    and SnapMap slides as ``[{"title", "png", "notes"}]`` (None: none).
    ``results`` is the ``resultspages.Results`` of the Quantification slides.
    ``glance`` is the list of ``(label, value)`` facts of "At a glance" and
    ``timing`` the ``glance.Timing`` of "Timing" (None: left out)."""
    if spec is None:
        spec = reportspec.spec_from_sections(sections, "deck")
    items = reportspec.active(spec)
    sha = reportspec.option(spec, "sha")
    title = (details.get("title") or "").strip() or "Experiment report"
    deck = _Deck(title, reportspec.cover_of(spec)["accent"])
    ids = [sid for sid, _skip in items]
    methods = (details.get("methods") or "").strip()
    cal = (details.get("calibration") or "").strip()
    # the calibration statement rides on the summary slide when both are in
    # (as it always did), else it gets a slide of its own; never twice
    cal_text = cal if cal and not ("methods" in ids and cal in methods) else ""
    for sid, skip in items:
        deck.section = sid
        if sid == "cover":
            art = covers.art(reportspec.cover_of(spec), "pptx", cover_data)
            if art.note and notes is not None:
                notes.append(art.note)
            deck.note(_title_slide(deck, details, logo, art), sid,
                      kind="cover")
        elif sid == "contents":
            deck.contents_at = len(deck.slides)      # made in ``_assemble``
        elif sid == "glance" and glance:
            _glance_slides(deck, glance)
        elif sid == "summary":
            paras = paragraphs(details.get("summary"))
            if cal_text and "calibration" in ids:
                paras.append("Energy calibration: " + cal_text)
            _text_slides(deck, "Summary", paras, size=16)
        elif sid == "results":
            _results_slides(deck, results, skip)
        elif sid == "methods":
            _text_slides(deck, "Methods", paragraphs(methods), size=14)
        elif sid == "timing" and timing:
            _timing_slides(deck, timing)
        elif sid == "calibration":
            if cal_text and "summary" not in ids:
                _text_slides(deck, "Energy calibration", [cal_text])
        elif sid == "files" and file_rows:
            _files_slides(deck, file_rows, sha)
        elif sid == "metadata":
            _metadata_slides(deck, docs, skip)
        elif sid == "images" and image_pages is not None:
            _image_slides(deck, image_pages())
        elif sid == "figures" and figures:
            chosen = [f for i, f in enumerate(figures, 1)
                      if reportspec.figure_id(f, i) not in skip]
            _figure_slides(deck, chosen, render_images)
    if not deck.slides:
        raise PptxError("Nothing to put in the presentation: choose at least "
                        "one section that has content.")
    _assemble(deck, reportspec.option(spec, "dividers"))
    cp = deck.prs.core_properties
    cp.title = title
    cp.author = (details.get("operator") or "").strip() or appinfo.NAME
    deck.prs.save(path)
    return len(deck.prs.slides)
