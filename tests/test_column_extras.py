"""The extras of the column-text importer: the SPECS ``.xy`` export (header
facts, kinetic axis checked against the header), Excel workbooks (sheets, the
dialog's sheet choice, ``openpyxl`` missing), the wide CSV this program writes
(read back spectrum by spectrum) and the photon / pass energy of each file in
the import dialog.

``XPS_SPECS_XY=<file or folder>`` also reads real SPECS ``.xy`` files and checks
their header against what was read.

Run:  python -m unittest discover tests
"""

import math
import os
import sys
import tempfile
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "tests"))

import columntext as ct  # noqa: E402
import exporters  # noqa: E402
import readers  # noqa: E402
from readers import column_text  # noqa: E402
from readers.base import Region  # noqa: E402
from test_column_text import peak, text  # noqa: E402

try:
    import openpyxl
    HAVE_XLSX = True
except Exception:                                    # pragma: no cover
    openpyxl, HAVE_XLSX = None, False

try:
    import tkinter as tk
    import spectradeck as ee
    HAVE_APP = ee.HAVE_MPL
except Exception:                                    # pragma: no cover
    tk, ee, HAVE_APP = None, None, False


def specs_text(ke, counts, hv=1253.6, pe=300, dwell=2, scans=1, values=None,
               extra=(), eol="\r\n"):
    """A SpecsLab Prodigy ``.xy`` file in the layout of a real one (20 header
    lines, a ``@`` line, then ``KE counts`` rows)."""
    head = [
        "% Group: Group1", "%", "% Region:            Region1",
        "% Anylsis Method:    XPS",
        "% Analyzer:          PHOIBOS HSA3500 100 R5 [HWType 7] MCD-5",
        "% Analyzer Lens:     MediumArea",
        "% Analyzer Slit:     6:3dia.\\2:open",
        "% Scan Mode:         FixedAnalyzerTransmission",
        f"% Number of Scans:   {scans}", "% Curves/Scan:       1",
        f"% Values/Curve:      {values if values is not None else len(ke)}",
        f"% Dwell Time:        {dwell}", f"% Excitation Energy: {hv}",
        f"% Kinetic Energy:    {ke[0]:g}", f"% Pass Energy:       {pe}",
        "% Bias Voltage:      100", "% Detector Voltage:  2800",
        "% Eff. Workfunction: 4.33", "% Source:            XRayDummy",
        "% Comment:           ", "%", "% Cycle: 0, Curve: 0", "%", "@"]
    rows = [f"{k:g} {c:g}" for k, c in zip(ke, counts)]
    return eol.join(head + rows + list(extra)) + eol


def ke_peak(n=80, hv=1253.6, be_peak=285.0):
    ke = [900.0 + i for i in range(n)]
    return ke, [round(50 + 900 * math.exp(-((hv - k - be_peak) / 3.0) ** 2), 2)
                for k in ke]


class Base(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dir = tempfile.mkdtemp()

    @classmethod
    def tearDownClass(cls):
        import shutil
        shutil.rmtree(cls.dir, ignore_errors=True)

    def write(self, name, data):
        p = os.path.join(self.dir, name)
        mode = "wb" if isinstance(data, bytes) else "w"
        with open(p, mode, **({} if mode == "wb" else
                              {"encoding": "utf-8", "newline": ""})) as fh:
            fh.write(data)
        return p


class TestSpecs(Base):
    def test_the_header_gives_the_energies_and_the_axis(self):
        ke, y = ke_peak()
        p = self.write("a.xy", specs_text(ke, y))
        self.assertIs(readers.reader_for(p), column_text.ColumnTextFile)
        d = readers.load_file(p)
        r = d.regions[0]
        self.assertEqual((r.photon_energy, r.pass_energy, r.dwell),
                         (1253.6, 300.0, 2.0))
        self.assertEqual(r.extra["n_scans"], 1)
        self.assertEqual(r.extra["analyser_mode"],
                         "Constant analyser energy (CAE)")
        self.assertEqual(r.lens_mode, "MediumArea")
        self.assertEqual(r.anode, "Mg Kα")
        self.assertEqual(r.count_units, "a.u.")          # the file gives none
        self.assertEqual(r.energy_label, "Binding Energy")
        self.assertAlmostEqual(r.energy[0], 1253.6 - ke[0], places=6)
        self.assertGreater(r.energy[0], r.energy[-1])     # BE high to low
        self.assertEqual(d.instrument["Work function (eV)"], "4.33")
        self.assertIn("PHOIBOS", d.instrument["Instrument"])
        self.assertEqual(d.warnings, [])

    def test_the_peak_lands_where_the_header_says(self):
        ke, y = ke_peak(be_peak=285.0)
        r = readers.load_file(self.write("b.xy", specs_text(ke, y))).regions[0]
        top = r.energy[max(range(len(y)), key=lambda i: r.counts[i])]
        self.assertAlmostEqual(top, 285.0, delta=1.0)

    def test_a_bare_xy_is_a_diffraction_pattern_not_a_spectrum(self):
        two_theta = [10 + 0.02 * i for i in range(200)]
        p = self.write("xrd.xy", text(zip(two_theta, range(200)), sep=" "))
        with self.assertRaises(readers.UnsupportedFormat):
            readers.reader_for(p)

    def test_the_value_count_in_the_header_is_checked(self):
        ke, y = ke_peak()
        p = self.write("c.xy", specs_text(ke, y, values=len(ke) + 5))
        self.assertTrue(any("5" in w and "values per curve" in w
                            for w in readers.load_file(p).warnings))

    def test_the_kinetic_axis_needs_the_headers_own_start_value(self):
        ke, y = ke_peak()
        text_ = specs_text(ke, y).replace(f"% Kinetic Energy:    {ke[0]:g}",
                                          "% Kinetic Energy:    5")
        t = ct.load_table(self.write("d.xy", text_))
        self.assertEqual(ct.guess_options(t)["axis"], "BE")   # not assumed

    def test_a_second_block_is_said_not_silently_dropped(self):
        ke, y = ke_peak()
        tail = ["%", "% Cycle: 0, Curve: 1", "%", "@", "abc def"]
        p = self.write("e.xy", specs_text(ke, y, extra=tail))
        self.assertTrue(any("several blocks" in w
                            for w in readers.load_file(p).warnings))

    @unittest.skipUnless(os.environ.get("XPS_SPECS_XY"), "set XPS_SPECS_XY")
    def test_real_files(self):
        root = os.environ["XPS_SPECS_XY"]
        paths = ([os.path.join(root, f) for f in os.listdir(root)
                  if f.lower().endswith(".xy")] if os.path.isdir(root)
                 else [root])
        for p in paths:
            d = readers.load_file(p)
            t = ct.load_table(p)
            r = d.regions[0]
            self.assertEqual(r.photon_energy,
                             ct.fact_float(t.facts, "excitation energy"), p)
            self.assertEqual(len(r.energy),
                             int(ct.fact_float(t.facts, "values/curve")), p)


class TestWide(Base):
    def regions(self):
        be, y = peak(60, 280, 292, 285)
        o_be, o_y = peak(40, 525, 540, 531)
        out = []
        for k, (name, e, c, sample) in enumerate((("C 1s", be, y, "S1"),
                                                  ("O 1s", o_be, o_y, "S1"),
                                                  ("Survey", o_be[:30],
                                                   o_y[:30], "Pt foil 2"))):
            out.append(Region(name=name, index=k, offset=k + 1, energy=list(e),
                              counts=list(c), decodable=True, sample=sample,
                              count_units="counts/s", photon_energy=1486.6))
        return out

    def test_the_csv_this_program_writes_comes_back_spectrum_by_spectrum(self):
        regs = self.regions()
        p = os.path.join(self.dir, "wide.csv")
        exporters.export_csv(regs, p)
        self.assertIs(readers.reader_for(p), column_text.ColumnTextFile)
        d = readers.load_file(p)
        self.assertEqual(len(d.regions), 3)
        for a, b in zip(regs, d.regions):
            self.assertEqual((b.name, b.sample, b.count_units, b.energy_label),
                             (a.name, a.sample, "counts/s", "Binding Energy"))
            self.assertEqual(b.energy, a.energy)
            self.assertEqual(b.counts, a.counts)
        self.assertEqual(len(d.regions[2].energy), 30)    # its own length

    def test_the_fit_columns_beside_them_are_not_spectra(self):
        regs = self.regions()
        p = os.path.join(self.dir, "wide2.csv")
        exporters.export_csv(regs[:1], p)
        with open(p, encoding="utf-8-sig") as fh:
            lines = fh.read().splitlines()
        lines[0] += ",S1 C 1s fit: background,S1 C 1s fit: envelope"
        lines[1:] = [ln + ",1.5,2.5" for ln in lines[1:]]
        with open(p, "w", encoding="utf-8-sig", newline="") as fh:
            fh.write("\r\n".join(lines) + "\r\n")
        d = readers.load_file(p)
        self.assertEqual([r.name for r in d.regions], ["C 1s"])

    def test_a_kinetic_export_needs_the_photon_energy_given(self):
        regs = self.regions()
        for r in regs:
            r.energy_label = "Kinetic Energy"
            r.energy = r.energy[::-1]
            r.counts = r.counts[::-1]
        p = os.path.join(self.dir, "ke.csv")
        exporters.export_csv(regs[:2], p)
        d = readers.load_file(p, options={"photon_energy": 1486.6})
        self.assertEqual({r.energy_label for r in d.regions}, {"Binding Energy"})
        self.assertAlmostEqual(d.regions[0].energy[0],
                               1486.6 - regs[0].energy[0], places=6)
        d = readers.load_file(p)
        self.assertEqual({r.energy_label for r in d.regions}, {"Kinetic Energy"})
        self.assertTrue(d.warnings)

    def test_the_head_of_a_column_is_split_into_sample_and_name(self):
        self.assertEqual(ct.split_head("S1 C 1s"), ("S1", "C 1s"))
        self.assertEqual(ct.split_head("Depth 3 Pt 4f loss"),
                         ("Depth 3", "Pt 4f loss"))
        self.assertEqual(ct.split_head("Pt foil 2 Survey"),
                         ("Pt foil 2", "Survey"))
        self.assertEqual(ct.split_head("Mystery"), ("", "Mystery"))

    def test_a_plain_table_is_not_taken_for_a_wide_one(self):
        be, y = peak()
        self.assertIsNone(ct.read_wide(text(zip(be, y),
                                            ["Binding Energy (eV)", "CPS"])))


@unittest.skipUnless(HAVE_XLSX, "openpyxl not installed")
class TestExcel(Base):
    def book(self, name, sheets):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for title, rows in sheets:
            ws = wb.create_sheet(title)
            for row in rows:
                ws.append(list(row))
        p = os.path.join(self.dir, name)
        wb.save(p)
        return p

    def test_a_sheet_with_a_header_and_a_table_reads_like_a_text_file(self):
        be, y = peak()
        rows = [["Sample S9"], ["Binding Energy (eV)", "CPS_C 1s"]] \
            + [[a, b] for a, b in zip(be, y)]
        p = self.book("one.xlsx", [("Data", rows)])
        self.assertIs(readers.reader_for(p), column_text.ColumnTextFile)
        d = readers.load_file(p)
        r = d.regions[0]
        self.assertEqual(r.counts, y)
        self.assertEqual(r.energy, be)
        self.assertEqual(r.count_units, "counts/s")

    def test_the_longest_table_wins_and_the_others_are_named(self):
        be, y = peak(60)
        o_be, o_y = peak(20, 525, 540, 531)
        p = self.book("two.xlsx", [
            ("Notes", [["Instrument", "Quantes"], ["Operator", "DM"]]),
            ("O 1s", [[a, b] for a, b in zip(o_be, o_y)]),
            ("C 1s", [[a, b] for a, b in zip(be, y)])])
        d = readers.load_file(p)
        self.assertEqual(len(d.regions[0].energy), 60)
        self.assertTrue(any("O 1s" in w and "were not" in w
                            for w in d.warnings))
        d = readers.load_file(p, options={"sheet": "O 1s"})
        self.assertEqual(len(d.regions[0].energy), 20)
        self.assertEqual(d.import_options["sheet"], "O 1s")

    def test_a_workbook_of_the_wide_export_comes_back_as_its_spectra(self):
        regs = TestWide().regions()
        p = os.path.join(self.dir, "wide.csv")
        exporters.export_csv(regs, p)
        import csv
        with open(p, encoding="utf-8-sig", newline="") as fh:
            rows = [[float(c) if c.replace(".", "", 1).replace("-", "", 1)
                     .isdigit() else (c if c else None) for c in row]
                    for row in csv.reader(fh)]
        x = self.book("wide.xlsx", [("Spectra", rows)])
        d = readers.load_file(x)
        self.assertEqual([r.name for r in d.regions], ["C 1s", "O 1s", "Survey"])
        self.assertEqual(d.regions[1].counts, regs[1].counts)

    def test_a_workbook_without_a_table_is_not_a_spectrum(self):
        p = self.book("none.xlsx", [("A", [["a", "b"], [1, 2]])])
        self.assertIsNone(ct.load_table(p))
        with self.assertRaises(ValueError):
            readers.load_file(p)

    def test_a_text_file_called_xlsx_is_not_a_workbook(self):
        be, y = peak()
        p = self.write("fake.xlsx", text(zip(be, y)))
        with self.assertRaises(readers.UnsupportedFormat):
            readers.reader_for(p)

    def test_without_openpyxl_it_says_so(self):
        be, y = peak()
        p = self.book("x.xlsx", [("Data", [[a, b] for a, b in zip(be, y)])])
        saved = sys.modules.get("openpyxl")
        sys.modules["openpyxl"] = None              # import raises ImportError
        try:
            with self.assertRaises(ct.WorkbookUnavailable) as cm:
                ct.load_table(p)
            self.assertIn("openpyxl", str(cm.exception))
        finally:
            sys.modules["openpyxl"] = saved


def items_for(paths):
    out = []
    for p in paths:
        t = ct.load_table(p)
        out.append((p, t, ct.guess_options(t, p)))
    return out


@unittest.skipUnless(HAVE_APP, "matplotlib / Tk not available")
class TestDialog(Base):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        import matplotlib
        cls._rc = matplotlib.rcParams.copy()
        try:
            cls.root = tk.Tk()
        except tk.TclError:
            raise unittest.SkipTest("no display")
        cls.root.withdraw()
        cls.ws = ee.Workspace(cls.root)
        be, y = peak(60, 280, 292, 285)
        ke, yk = ke_peak()
        cls.plain = cls.write_(cls, "plain.asc", text(zip(be, y),
                                                      ["BE_C1s", "CPS_C1s"]))
        cls.specs = cls.write_(cls, "s.xy", specs_text(ke, yk, hv=1253.6,
                                                       pe=300))

    @staticmethod
    def write_(cls, name, data):
        p = os.path.join(cls.dir, name)
        with open(p, "w", encoding="utf-8", newline="") as fh:
            fh.write(data)
        return p

    @classmethod
    def tearDownClass(cls):
        import matplotlib
        matplotlib.rcParams.update(cls._rc)
        cls.root.destroy()
        super().tearDownClass()

    def dialog(self, paths, defaults=None):
        import columnimport_ui
        dlg = columnimport_ui.ColumnImportDialog(
            self.root, self.ws, items_for(paths), defaults)
        self.addCleanup(lambda: dlg.winfo_exists() and dlg.destroy())
        return dlg

    def test_each_file_has_its_own_photon_and_pass_energy(self):
        dlg = self.dialog([self.plain, self.specs])
        dlg.set_row(0, photon="Al Kα (1486.6 eV)", pass_energy="20")
        dlg.accept()
        a, b = dlg.result[self.plain], dlg.result[self.specs]
        self.assertEqual((a["photon_energy"], a["pass_energy"], a["anode"]),
                         (1486.6, 20.0, "Al Kα"))
        self.assertEqual((b["photon_energy"], b["pass_energy"], b["anode"]),
                         (1253.6, 300.0, "Mg Kα"))              # its header's
        self.assertEqual(b["axis"], "KE")

    def test_the_header_values_start_in_the_table_and_the_default_fills_the_rest(self):
        dlg = self.dialog([self.plain, self.specs],
                          {"photon_energy": 1486.6, "pass_energy": 40})
        shown = [dlg.tree.item(str(i))["values"][5] for i in range(2)]
        self.assertEqual(shown, ["Al Kα (1486.6 eV)", "Mg Kα (1253.6 eV)"])
        self.assertEqual(dlg.rows[0]["opts"]["pass_energy"], 40.0)
        self.assertEqual(dlg.rows[1]["opts"]["pass_energy"], 300.0)

    def test_the_every_file_boxes_set_all_files_at_once(self):
        dlg = self.dialog([self.plain, self.specs])
        dlg.set_shared(photon="Unknown")
        for row in dlg.rows:
            self.assertIsNone(row["opts"]["photon_energy"])
        dlg.set_shared(photon="1486.6", pass_energy=10)
        for row in dlg.rows:
            self.assertEqual((row["opts"]["photon_energy"],
                              row["opts"]["pass_energy"]), (1486.6, 10.0))

    def test_the_selected_file_boxes_follow_the_selection(self):
        dlg = self.dialog([self.plain, self.specs])
        dlg.tree.selection_set("1")
        dlg._on_select()
        self.assertEqual(dlg.v_hv.get(), "Mg Kα (1253.6 eV)")
        self.assertEqual(dlg.v_pe.get(), "300")
        dlg.v_hv.set("1486.6")
        dlg._commit_energies()
        self.assertEqual(dlg.rows[1]["opts"]["photon_energy"], 1486.6)
        self.assertIsNone(dlg.rows[0]["opts"]["photon_energy"])

    def test_a_wide_file_shows_its_spectra_and_locks_its_columns(self):
        regs = TestWide().regions()
        p = os.path.join(self.dir, "w.csv")
        exporters.export_csv(regs, p)
        dlg = self.dialog([p])
        self.assertEqual(dlg.tree.item("0")["values"][6], "3 spectra")
        dlg.tree.selection_set("0")
        dlg._on_select()
        self.assertEqual(str(dlg.cb_energy.cget("state")), "disabled")
        dlg.set_row(0, axis="KE", name="Nope")       # ignored: the header rules
        self.assertEqual(dlg.rows[0]["opts"]["axis"], "BE")
        dlg.accept()
        self.assertEqual(readers.load_file(p, options=dlg.result[p])
                         .regions[1].name, "O 1s")

    @unittest.skipUnless(HAVE_XLSX, "openpyxl not installed")
    def test_the_sheet_can_be_chosen(self):
        be, y = peak(60)
        o_be, o_y = peak(20, 525, 540, 531)
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for title, rows in (("O 1s", list(zip(o_be, o_y))),
                            ("C 1s", list(zip(be, y)))):
            ws = wb.create_sheet(title)
            for r in rows:
                ws.append(list(r))
        p = os.path.join(self.dir, "sheets.xlsx")
        wb.save(p)
        dlg = self.dialog([p])
        self.assertEqual(dlg.rows[0]["table"].sheet, "C 1s")
        dlg.set_row(0, photon="1486.6")
        dlg.set_sheet(0, "O 1s")
        self.assertEqual(dlg.rows[0]["table"].n_rows, 20)
        self.assertEqual(dlg.rows[0]["opts"]["photon_energy"], 1486.6)
        self.assertEqual(dlg.rows[0]["opts"]["sheet"], "O 1s")
        dlg.accept()
        self.assertEqual(dlg.result[p]["sheet"], "O 1s")


@unittest.skipUnless(HAVE_APP and HAVE_XLSX, "Tk, matplotlib or openpyxl missing")
class TestOpenFolder(Base):
    """A folder holding a workbook with a spectrum, one without, and a SPECS
    file: the first and the last open (through the import dialog), the middle
    one is skipped like any file that is not a spectrum."""

    def test_only_workbooks_with_a_table_are_opened(self):
        try:
            root = tk.Tk()
        except tk.TclError:
            self.skipTest("no display")
        root.withdraw()
        import matplotlib
        rc = matplotlib.rcParams.copy()
        boxes = (ee.messagebox.showinfo, ee.messagebox.showwarning,
                 ee.messagebox.showerror)
        ee.messagebox.showinfo = ee.messagebox.showwarning =             ee.messagebox.showerror = lambda *a, **k: None
        try:
            ws = ee.Workspace(root)
            folder = os.path.join(self.dir, "mixed")
            os.makedirs(folder, exist_ok=True)
            be, y = peak()
            wb = openpyxl.Workbook()
            for r in zip(be, y):
                wb.active.append(list(r))
            wb.save(os.path.join(folder, "spec.xlsx"))
            wb = openpyxl.Workbook()
            wb.active.append(["just", "words"])
            wb.save(os.path.join(folder, "report.xlsx"))
            ke, yk = ke_peak()
            with open(os.path.join(folder, "s.xy"), "w", newline="") as fh:
                fh.write(specs_text(ke, yk))
            asked = []

            def answer(items):
                asked.append([os.path.basename(p) for p, _t, _o in items])
                return {p: dict(o, sample="S1") for p, _t, o in items}
            ws._ask_column_import = answer
            ws._open_folder_path(folder)
            self.assertEqual(sorted(asked[0]), ["s.xy", "spec.xlsx"])
            self.assertEqual(sorted(os.path.basename(p.path)
                                    for p in ws.docs), ["s.xy", "spec.xlsx"])
            self.assertIn("1 other file", ws._folder_note)
            xy = next(p for p in ws.docs if p.path.endswith(".xy"))
            self.assertEqual(xy.regions[0].photon_energy, 1253.6)
        finally:
            ee.messagebox.showinfo, ee.messagebox.showwarning,                 ee.messagebox.showerror = boxes
            matplotlib.rcParams.update(rc)
            root.destroy()


if __name__ == "__main__":
    unittest.main()
