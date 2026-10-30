"""The repeat statistics of the REELS π / (σ+π) table (reelstable.py): the mean,
sample standard deviation (n - 1) and n of the ratio, the sp² fraction and the
peak positions over the groups of samples that repeat one another, and where
they go (the dialog's lines, the CSV, the PDF, the slides, Word, the hand-over).

Run:  python -m unittest discover tests
"""

import csv
import io
import os
import sys
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
sys.path.insert(0, HERE)

import handover  # noqa: E402
import reelstable  # noqa: E402
import repeats  # noqa: E402
import reportspec as rs  # noqa: E402
from test_glance import HAVE_DOCX, HAVE_PDF, HAVE_PPTX, mupdf  # noqa: E402
from test_reelstable import Base, HOPG, item, stored  # noqa: E402


def keyed(sample, ratio, spectrum="REELS", doc="f", **over):
    """An item whose sample is a repeat-group member (key ``doc/sample``) with
    a chosen ratio (the areas are made consistent with it)."""
    kw = dict(ratio=ratio, i_pi=ratio * 1000.0, i_sp=1000.0,
              pos_pi=6.0 + ratio, pos_sp=26.0)
    kw.update(over)
    res = stored(**kw)
    it = item(sample, spectrum, res=res)
    it["key"] = f"{doc}/{sample}"
    return it


GROUP = {"Series": {"kind": "locations", "samples": ["f/A", "f/B", "f/C"]}}


class TestRepeatGroups(unittest.TestCase):
    def table(self, **kw):
        items = kw.pop("items", [keyed("A", 0.10), keyed("B", 0.12),
                                 keyed("C", 0.14)])
        return reelstable.collect(items, HOPG, kw.pop("groups", GROUP))

    def test_mean_sd_and_n_of_the_ratio(self):
        t = self.table()
        self.assertEqual(len(t.groups), 1)
        g = t.groups[0]
        self.assertEqual((g.name, g.kind, g.members),
                         ("Series", "locations", ["A", "B", "C"]))
        ratio = next(r for r in g.rows if r.name == "Iπ/Iσ+π")
        self.assertAlmostEqual(ratio.mean, 0.12, 12)
        self.assertAlmostEqual(ratio.sd, 0.02, 12)          # n - 1
        self.assertEqual(ratio.n, 3)
        self.assertEqual([lab for lab, _v in ratio.values], ["A", "B", "C"])
        self.assertEqual([round(v, 2) for _l, v in ratio.values],
                         [0.10, 0.12, 0.14])

    def test_the_other_quantities_are_there(self):
        g = self.table().groups[0]
        self.assertEqual([r.name for r in g.rows],
                         ["Iπ/Iσ+π", "sp² fraction", "π position (eV)",
                          "σ+π position (eV)"])
        self.assertAlmostEqual(g.rows[1].mean, 0.12 / 0.16, 9)   # HOPG = 1
        self.assertAlmostEqual(g.rows[3].sd, 0.0, 12)            # all 26.0

    def test_a_missing_value_is_left_out_of_that_quantity_only(self):
        items = [keyed("A", 0.10), keyed("B", 0.12),
                 keyed("C", 0.14, pos_pi=None)]
        g = self.table(items=items).groups[0]
        self.assertEqual(next(r for r in g.rows
                              if r.name == "π position (eV)").n, 2)
        self.assertEqual(next(r for r in g.rows
                              if r.name == "Iπ/Iσ+π").n, 3)

    def test_one_spectrum_has_a_mean_and_no_sd_and_says_so(self):
        g = self.table(items=[keyed("A", 0.10)],
                       groups={"Solo": {"kind": "repeats",
                                        "samples": ["f/A"]}}).groups[0]
        self.assertIsNone(g.rows[0].sd)
        self.assertEqual(g.rows[0].n, 1)
        self.assertTrue(any("Only one spectrum" in n for n in g.notes))

    def test_a_group_with_no_reels_spectrum_is_about_something_else(self):
        t = self.table(groups={"Fits": {"kind": "samples",
                                        "samples": ["f/X", "f/Y"]}})
        self.assertEqual(t.groups, [])
        self.assertEqual(reelstable.collect([keyed("A", 0.1)], HOPG,
                                            None).groups, [])
        self.assertEqual(reelstable.collect([keyed("A", 0.1)], HOPG,
                                            {}).groups, [])

    def test_a_member_with_no_ratio_is_named_and_left_out(self):
        g = self.table(groups={"G": {"kind": "locations",
                                     "samples": ["f/A", "f/B", "f/Z"]}}
                       ).groups[0]
        self.assertEqual(g.members, ["A", "B"])
        self.assertTrue(any("Z: no REELS ratio measured" in n
                            for n in g.notes))
        self.assertEqual(g.rows[0].n, 2)

    def test_other_windows_mean_no_statistics(self):
        items = [keyed("A", 0.10), keyed("B", 0.12),
                 keyed("C", 0.14, pi_window=[3.0, 8.0])]
        g = self.table(items=items).groups[0]
        self.assertEqual(g.rows, [])
        self.assertEqual(len(g.members), 3)
        self.assertTrue(any("not comparable" in n for n in g.notes))

    def test_a_flagged_spectrum_is_counted_and_named(self):
        items = [keyed("A", 0.10), keyed("B", 0.12),
                 keyed("C", 0.14, dip_pi=0.5)]
        g = self.table(items=items).groups[0]
        self.assertEqual(g.rows[0].n, 3)
        self.assertTrue(any("flagged" in n and "C" in n for n in g.notes))

    def test_a_sample_with_two_spectra_gives_two_values_labelled(self):
        items = [keyed("A", 0.10), keyed("A", 0.12, spectrum="REELS 2"),
                 keyed("B", 0.14)]
        g = self.table(items=items).groups[0]
        self.assertEqual(g.rows[0].n, 3)
        self.assertEqual([lab for lab, _v in g.rows[0].values],
                         ["A / REELS", "A / REELS 2", "B"])

    def test_kind_text_and_several_groups(self):
        groups = {"One": {"kind": "locations", "samples": ["f/A", "f/B"]},
                  "Two": {"kind": "repeats", "samples": ["f/C", "f/B"]}}
        t = self.table(groups=groups)
        self.assertEqual([g.name for g in t.groups], ["One", "Two"])
        self.assertIn("locations", t.groups[0].kind_text)
        self.assertIn("repeat", t.groups[1].kind_text)

    def test_lines_for_the_dialog(self):
        lines = reelstable.group_lines(self.table())
        self.assertEqual(len(lines), 1)
        self.assertIn("Series (3 spectra)", lines[0])
        self.assertIn("ratio 0.1200 ± 0.0200", lines[0])
        self.assertIn("sp² 0.75", lines[0])
        none = reelstable.group_lines(self.table(
            items=[keyed("A", 0.1), keyed("B", 0.1, pi_window=[3, 8.5])],
            groups={"G": {"kind": "locations", "samples": ["f/A", "f/B"]}}))
        self.assertIn("no statistics", none[0])

    def test_the_report_rows(self):
        rows = reelstable.repeat_report_rows(self.table().groups[0])
        self.assertEqual(rows[0][:4], ["Iπ/Iσ+π", "0.1200", "0.0200", "3"])
        self.assertIn("A=0.1000", rows[0][4])
        self.assertEqual(len(rows[0]), len(reelstable.repeat_header()))
        solo = self.table(items=[keyed("A", 0.1)], groups={
            "S": {"kind": "repeats", "samples": ["f/A"]}}).groups[0]
        self.assertEqual(reelstable.repeat_report_rows(solo)[0][2],
                         reelstable.DASH)

    def test_the_csv(self):
        t = self.table()
        rows = reelstable.repeat_csv_rows(t)
        self.assertEqual(rows[0], ["Group", "Kind", "Quantity", "Mean",
                                   "SD (n-1)", "n", "Individual values"])
        self.assertEqual(rows[1][:3], ["Series", "different locations on the "
                                       "sample", "Iπ/Iσ+π"])
        self.assertEqual((rows[1][3], rows[1][4], rows[1][5]),
                         ("0.12", "0.02", "3"))
        self.assertIn("A=0.1", rows[1][6])
        data = reelstable.repeat_csv_bytes(t)
        self.assertTrue(data.startswith(b"\xef\xbb\xbf"))
        back = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
        self.assertEqual(back, rows)
        empty = self.table(groups={})
        self.assertEqual(reelstable.repeat_csv_rows(empty), [])
        self.assertEqual(reelstable.repeat_csv_bytes(empty), b"")

    def test_samples_for_the_group_suggestions(self):
        t = self.table(items=[keyed("Gunsan - location 1", 0.1),
                              keyed("Gunsan - location 2", 0.1),
                              keyed("Other", 0.1)])
        stand_ins = reelstable.suggest_samples(t)
        self.assertEqual([s.key for s in stand_ins],
                         ["f/Gunsan - location 1", "f/Gunsan - location 2",
                          "f/Other"])
        got = repeats.suggest(stand_ins)
        self.assertEqual(list(got), ["Gunsan"])
        self.assertEqual(len(got["Gunsan"]["samples"]), 2)

    def test_the_stats_use_the_tables_own_values(self):
        t = self.table()
        col = [r.ratio for r in t.rows]
        self.assertAlmostEqual(t.groups[0].rows[0].mean, sum(col) / 3, 12)


@unittest.skipUnless(HAVE_PDF, "reportlab / PyMuPDF / matplotlib missing")
class TestInThePdf(Base):
    def build(self, table):
        import report
        self.k = getattr(self, "k", 0) + 1
        path = os.path.join(self.dir, f"g{self.k}.pdf")
        report.build_report(path, self.details, "", self.rows, self.docs, [],
                            lambda *a: 0, spec=rs.default_spec(),
                            glance=self.facts, timing=self.timing,
                            reels=table)
        doc = mupdf.open(path)
        self.addCleanup(doc.close)
        return doc, "\n".join(p.get_text() for p in doc)

    def test_the_groups_are_in_the_pdf_and_the_bookmarks(self):
        table = reelstable.collect(
            [keyed("Film A", 0.10), keyed("Film B", 0.12),
             keyed("Film C", 0.14)], HOPG,
            {"Series": {"kind": "locations",
                        "samples": ["f/Film A", "f/Film B", "f/Film C"]}})
        doc, text = self.build(table)
        for token in ("Repeat groups", "Series", "SD (n−1)", "0.1200",
                      "0.0200", "nothing is re-measured"):
            self.assertIn(token, text)
        self.assertIn("REELS analysis", [t for _l, t, _p in doc.get_toc()])

    def test_no_groups_no_subsection(self):
        _doc, text = self.build(reelstable.collect(
            [keyed("Film A", 0.10)], HOPG, None))
        self.assertNotIn("Repeat groups", text)


@unittest.skipUnless(HAVE_PPTX, "python-pptx / PyMuPDF not installed")
class TestInTheDeck(Base):
    def build(self, table):
        import pptx_export
        from pptx import Presentation
        self.k = getattr(self, "k", 0) + 1
        path = os.path.join(self.dir, f"g{self.k}.pptx")
        pptx_export.build_deck(path, self.details, "", self.rows, self.docs,
                               [], lambda n, f: [], spec=rs.default_spec(),
                               glance=self.facts, timing=self.timing,
                               reels=table)
        return list(Presentation(path).slides)

    @staticmethod
    def title(slide):
        return slide.shapes.title.text if slide.shapes.title is not None \
            else ""

    def test_a_slide_per_group(self):
        slides = self.build(reelstable.collect(
            [keyed("A", 0.10), keyed("B", 0.12)], HOPG,
            {"Series": {"kind": "locations", "samples": ["f/A", "f/B"]},
             "Mixed": {"kind": "repeats", "samples": ["f/A", "f/B"]}}))
        titles = [self.title(s) for s in slides]
        self.assertIn("REELS repeat group – Series", titles)
        self.assertIn("REELS repeat group – Mixed", titles)
        sl = slides[titles.index("REELS repeat group – Series")]
        cells = [[c.text for c in row.cells] for sh in sl.shapes
                 if sh.has_table for row in sh.table.rows]
        self.assertEqual(cells[0], reelstable.repeat_header())
        self.assertEqual(cells[1][:4], ["Iπ/Iσ+π", "0.1100", "0.0141", "2"])

    def test_a_group_with_no_statistics_is_a_text_slide(self):
        slides = self.build(reelstable.collect(
            [keyed("A", 0.10), keyed("B", 0.12, pi_window=[3, 8.5])], HOPG,
            {"G": {"kind": "locations", "samples": ["f/A", "f/B"]}}))
        words = " ".join(
            sh.text_frame.text for sl in slides
            if self.title(sl).startswith("REELS repeat group")
            for sh in sl.shapes if sh.has_text_frame)
        self.assertIn("not comparable", words)


@unittest.skipUnless(HAVE_DOCX, "python-docx not installed")
class TestInWord(Base):
    def test_a_heading_and_a_table_per_group(self):
        import docx_export
        from docx import Document
        table = reelstable.collect(
            [keyed("A", 0.10), keyed("B", 0.12)], HOPG,
            {"Series": {"kind": "locations", "samples": ["f/A", "f/B"]}})
        path = os.path.join(self.dir, "g.docx")
        docx_export.build_document(
            path, self.details, "", self.rows, self.docs, [],
            lambda n, f: [], spec=rs.default_spec(), glance=self.facts,
            timing=self.timing, reels=table)
        doc = Document(path)
        heads = [p.text for p in doc.paragraphs
                 if p.style.name.startswith("Heading")]
        self.assertIn("Repeat groups", heads)
        tables = [[[c.text for c in row.cells] for row in t.rows]
                  for t in doc.tables]
        mine = next(t for t in tables if t[0] == reelstable.repeat_header())
        self.assertEqual(mine[1][:4], ["Iπ/Iσ+π", "0.1100", "0.0141", "2"])


class TestHandover(unittest.TestCase):
    def test_a_second_part_only_with_statistics(self):
        with_groups = reelstable.collect(
            [keyed("A", 0.10), keyed("B", 0.12)], HOPG,
            {"G": {"kind": "locations", "samples": ["f/A", "f/B"]}})
        parts = handover.reels_parts(with_groups)
        self.assertEqual([p.arc for p in parts],
                         [handover.REELS_ARC, handover.REELS_GROUPS_ARC])
        self.assertEqual(parts[1].data,
                         reelstable.repeat_csv_bytes(with_groups))
        self.assertIn("About the REELS repeat groups",
                      handover.readme_text({}, parts))
        none = handover.reels_parts(reelstable.collect(
            [keyed("A", 0.10)], HOPG, None))
        self.assertEqual([p.arc for p in none], [handover.REELS_ARC])
        self.assertNotIn("About the REELS repeat groups",
                         handover.readme_text({}, none))
        mixed = handover.reels_parts(reelstable.collect(
            [keyed("A", 0.10), keyed("B", 0.12, pi_window=[3, 8.5])], HOPG,
            {"G": {"kind": "locations", "samples": ["f/A", "f/B"]}}))
        self.assertEqual([p.arc for p in mixed], [handover.REELS_ARC])


if __name__ == "__main__":
    unittest.main()
