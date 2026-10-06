"""The Report generator's model and its effect on the PDF, the slides and the
workspace (milestone G2).

Run:  python -m unittest discover tests
"""

import io
import json
import os
import shutil
import sys
import tempfile
import unittest
from types import SimpleNamespace

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
sys.path.insert(0, HERE)

import reportspec as rs  # noqa: E402

try:
    import reportlab  # noqa: F401
    try:
        import pymupdf as mupdf
    except ImportError:
        import fitz as mupdf
    HAVE_PDF = True
except ImportError:
    HAVE_PDF = False
try:
    import pptx  # noqa: F401
    import matplotlib  # noqa: F401
    HAVE_PPTX = True
except ImportError:
    HAVE_PPTX = False
try:
    import docx  # noqa: F401
    HAVE_DOCX = True
except ImportError:
    HAVE_DOCX = False


class TestSpec(unittest.TestCase):
    def test_the_default_is_everything_in_the_results_first_order(self):
        s = rs.default_spec()
        self.assertEqual(rs.order(s), list(rs.SECTION_IDS))
        self.assertEqual(rs.order(s)[:5], ["cover", "contents", "glance",
                                           "summary", "results"])
        self.assertEqual(rs.order(s).index("timing"),
                         rs.order(s).index("methods") + 1)
        self.assertEqual(rs.order(s)[-2:], ["metadata", "files"])   # appendix
        self.assertTrue(all(rs.is_on(s, i) for i in rs.SECTION_IDS))

    def test_the_cover_font_is_validated_and_kept_by_presets(self):
        import fonts
        self.assertEqual(rs.cover_of(rs.default_spec())["font"], fonts.FAMILY)
        for bad in (None, 3, "", "Comic Sans", "IBM Plex Mono", ["Inter"]):
            s = rs.sanitise({"cover": {"font": bad}})
            self.assertEqual(rs.cover_of(s)["font"], fonts.FAMILY, bad)
        s = rs.with_cover(rs.default_spec(), font="Inter")
        self.assertEqual(rs.cover_of(s)["font"], "Inter")
        # a spec saved before the font existed reads as the default
        old = rs.default_spec()
        del old["cover"]["font"]
        self.assertEqual(rs.cover_of(old)["font"], fonts.FAMILY)
        # a preset says what goes in, never the look
        for name, preset in rs.BUILTIN_PRESETS.items():
            got = rs.with_cover_of(preset, s)
            self.assertEqual(rs.cover_of(got)["font"], "Inter", name)

    def test_every_section_has_a_label_a_hint_and_a_short_name(self):
        # the PDF preview's toolbar indexes SHORT by section id
        self.assertEqual(set(rs.SHORT), set(rs.SECTION_IDS))
        self.assertEqual(set(rs.LABELS), set(rs.SECTION_IDS))
        self.assertEqual(set(rs.HINTS), set(rs.SECTION_IDS))
        self.assertEqual(set(rs.LEGACY_ORDER), set(rs.SECTION_IDS))

    def test_sanitise_repairs_anything(self):
        for bad in (None, 5, "x", [], {"sections": 3}, {"sections": [7, None]}):
            self.assertEqual(rs.sanitise(bad), rs.default_spec())
        s = rs.sanitise({"sections": [{"id": "files", "on": False},
                                      {"id": "nope", "on": True},
                                      {"id": "files", "on": True},
                                      {"id": "cover"}],
                         "skip": {"figures": ["a", "a", 3], "cover": ["x"],
                                  "metadata": []},
                         "options": {"sha": "weird", "other": 1}})
        self.assertEqual(rs.order(s)[:2], ["files", "cover"])     # kept order
        self.assertEqual(sorted(rs.order(s)), sorted(rs.SECTION_IDS))  # all there
        self.assertFalse(rs.is_on(s, "files"))
        self.assertTrue(rs.is_on(s, "summary"))       # a section the input lacked
        self.assertEqual(s["skip"], {"figures": ["a", "3"]})
        self.assertEqual(s["options"], {"sha": "short", "dividers": "auto",
                                     "mosaic": "on", "page": "a4", "rsf": "off"})

    def test_changes_return_new_specs_and_leave_the_old_alone(self):
        a = rs.default_spec()
        b = rs.with_on(a, "summary", False)
        self.assertTrue(rs.is_on(a, "summary"))
        self.assertFalse(rs.is_on(b, "summary"))
        self.assertEqual(rs.order(rs.moved(a, "figures", -5))[:3],
                         ["figures", "cover", "contents"])
        self.assertEqual(rs.order(rs.moved(a, "cover", -5)), rs.order(a))
        self.assertEqual(rs.order(rs.moved(a, "files", 9)), rs.order(a))
        self.assertEqual(rs.order(rs.moved(a, "nothing", 1)), rs.order(a))

    def test_children_are_on_unless_skipped(self):
        s = rs.with_child(rs.default_spec(), "figures", "fig2", False)
        self.assertEqual(rs.skipped(s, "figures"), {"fig2"})
        s = rs.with_child(s, "figures", "fig2", True)
        self.assertEqual(rs.skipped(s, "figures"), set())
        self.assertNotIn("figures", s["skip"])          # nothing left to store

    def test_select_all_and_none(self):
        s = rs.with_all(rs.with_child(rs.default_spec(), "figures", "x",
                                      False), False)
        self.assertEqual(rs.active(s), [])
        self.assertEqual(s["skip"], {})
        self.assertEqual(len(rs.active(rs.with_all(s, True))),
                         len(rs.SECTION_IDS))

    def test_active_lists_sections_in_order_with_their_skips(self):
        s = rs.with_child(rs.moved(rs.default_spec(), "files", -11),
                          "metadata",
                          "f1", False)
        got = rs.active(s)
        self.assertEqual([i for i, _k in got][0], "files")
        self.assertEqual(dict(got)["metadata"], {"f1"})
        self.assertEqual([i for i, _k in rs.active(s, {"cover", "files"})],
                         ["files", "cover"])

    def test_the_old_section_names_mean_what_they_did(self):
        pdf = rs.spec_from_sections(("cover", "figures"), "pdf")
        self.assertEqual([i for i, _ in rs.active(pdf)],
                         ["cover", "summary", "methods", "calibration",
                          "files", "figures"])
        deck = rs.spec_from_sections(("title", "metadata"), "deck")
        self.assertEqual([i for i, _ in rs.active(deck)],
                         ["cover", "summary", "methods", "calibration",
                          "metadata"])
        self.assertEqual(rs.active(rs.spec_from_sections((), "pdf")), [])

    def test_json_round_trip(self):
        s = rs.with_option(rs.with_child(rs.moved(
            rs.default_spec(), "images", -1), "figures", "f", False),
            "sha", "none")
        self.assertEqual(rs.sanitise(json.loads(json.dumps(s))), s)
        self.assertTrue(rs.same(s, json.loads(json.dumps(s))))
        self.assertFalse(rs.same(s, rs.default_spec()))

    def test_presets(self):
        allp = rs.all_presets({"Mine": rs.default_spec(),
                               "Everything": rs.with_all(rs.default_spec(),
                                                         False),
                               "": rs.default_spec()})
        self.assertEqual(list(allp)[:4], list(rs.BUILTIN_PRESETS))
        self.assertIn("Mine", allp)
        self.assertTrue(rs.is_on(allp["Everything"], "cover"))   # not shadowed
        self.assertEqual(rs.clean_presets(None), {})
        quick = rs.BUILTIN_PRESETS["Quick look"]
        self.assertEqual([i for i, _ in rs.active(quick)], ["cover", "figures"])
        audit = rs.BUILTIN_PRESETS["Audit trail"]
        self.assertNotIn("figures", [i for i, _ in rs.active(audit)])

    def test_the_presets_are_valid_specs(self):
        for name, spec in rs.BUILTIN_PRESETS.items():
            self.assertEqual(rs.sanitise(spec), spec, name)


class TestInventory(unittest.TestCase):
    def docs(self):
        return [SimpleNamespace(path="C:/x/a.vgd", file_id="f1"),
                SimpleNamespace(path="C:/x/b.vgd", file_id="")]

    def inv(self, **kw):
        args = dict(details={"summary": "s"}, methods_text="m",
                    calibration="c", file_rows=[{"name": "a"}],
                    docs=self.docs(),
                    figures=[{"id": "g1", "name": "One"}, {"name": "Two"}],
                    has_images=True, results=[("f1/S", "S")],
                    glance=[("Data", "1 file")], timing=["First start"])
        args.update(kw)
        return rs.inventory(**args)

    def test_what_is_there(self):
        inv = self.inv()
        self.assertEqual(inv.present_ids(), set(rs.SECTION_IDS))
        self.assertEqual(inv.children["figures"], [("g1", "One"),
                                                   ("fig2", "Two")])
        self.assertEqual(inv.children["metadata"], [("f1", "a.vgd"),
                                                    ("b.vgd", "b.vgd")])
        self.assertEqual(inv.summary("figures"), "2 figures")
        self.assertEqual(inv.summary("files"), "1 file")

    def test_what_is_missing_says_why(self):
        inv = self.inv(details={}, methods_text="", calibration=" ",
                       file_rows=[], docs=[], figures=[], has_images=False,
                       results=(), glance=[], timing=None)
        self.assertEqual(inv.present_ids(), {"cover", "contents"})
        self.assertIn("acquisition times", inv.summary("timing"))
        self.assertIn("summary", inv.summary("summary"))
        self.assertIn("Figures", inv.summary("figures"))
        self.assertIn("loaded", inv.summary("metadata"))

    def test_no_matplotlib_says_so(self):
        inv = self.inv(have_mpl=False)
        self.assertNotIn("figures", inv.present_ids())
        self.assertIn("matplotlib", inv.summary("figures"))

    def test_describe(self):
        inv = self.inv()
        s = rs.with_child(rs.default_spec(), "figures", "g1", False)
        self.assertIn("Figures (1 of 2)", rs.describe(s, inv))
        self.assertEqual(rs.describe(rs.with_all(s, False), inv),
                         "nothing selected")
        self.assertEqual(rs.describe(rs.BUILTIN_PRESETS["Quick look"], inv),
                         "Cover page, Figures")


class Tmp(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.dir, True)


@unittest.skipUnless(HAVE_PDF, "reportlab / PyMuPDF not installed")
class TestPdfFollowsTheSpec(Tmp):
    def setUp(self):
        super().setUp()
        from test_metasummary import Doc, region as mregion
        self.docs = [Doc([mregion("Survey", 160, step=1.0)], "a.vgd"),
                     Doc([mregion("Mo 3d", 40)], "b.vgd")]
        self.figs = [{"name": "One", "caption": "", "state": {}},
                     {"name": "Two", "caption": "", "state": {}},
                     {"name": "Three", "caption": "", "state": {}}]
        self.rows = [{"name": "a.vgd", "format": "Thermo", "regions": 1,
                      "size": 2048, "sha256": "ab" * 32}]
        self.details = {"title": "The Study", "customer": "ACME",
                        "summary": "Everything worked.",
                        "methods": "Spectra were recorded.",
                        "calibration": "C 1s at 284.8 eV."}
        self.calls = []

    def render(self, pdf, n, fig):
        from matplotlib.figure import Figure
        self.calls.append((n, fig["name"]))
        f = Figure(figsize=(11.7, 8.3))
        f.text(0.5, 0.5, f"figure {n} {fig['name']}")
        pdf.savefig(f)
        return 1

    def build(self, spec):
        import report
        path = os.path.join(self.dir, "r.pdf")
        report.build_report(path, self.details, "", self.rows, self.docs,
                            self.figs, self.render, spec=spec)
        with mupdf.open(path) as d:
            return [p.get_text() for p in d]

    def first_page_with(self, pages, token):
        return next(i for i, t in enumerate(pages) if token in t)

    def test_the_order_is_the_specs_order(self):
        # the customer is named on the cover only (the title is in every footer)
        pages = self.build(rs.default_spec())
        self.assertLess(self.first_page_with(pages, "ACME"),
                        self.first_page_with(pages, "figure 1"))
        self.assertLess(self.first_page_with(pages, "figure 3"),
                        self.first_page_with(pages, "Acquisition metadata - a.vgd"))
        first = self.build(rs.moved(rs.default_spec(), "figures", -5))
        self.assertIn("figure 1", first[0])              # figures lead now
        self.assertLess(self.first_page_with(first, "figure 1"),
                        self.first_page_with(first, "ACME"))

    def test_off_sections_are_left_out(self):
        spec = rs.with_on(rs.with_on(rs.default_spec(), "summary", False),
                          "metadata", False)
        text = "\n".join(self.build(spec))
        self.assertNotIn("Everything worked.", text)
        self.assertNotIn("Acquisition metadata", text)
        self.assertIn("Spectra were recorded.", text)

    def test_a_skipped_figure_is_not_drawn_and_the_rest_are_renumbered(self):
        spec = rs.with_child(rs.default_spec(), "figures", "fig1", False)
        self.build(spec)
        self.assertEqual(self.calls, [(1, "Two"), (2, "Three")])

    def test_a_skipped_file_has_no_metadata_pages(self):
        spec = rs.with_child(rs.default_spec(), "metadata", "b.vgd", False)
        text = "\n".join(self.build(spec))
        self.assertIn("Acquisition metadata - a.vgd", text)
        self.assertNotIn("Acquisition metadata - b.vgd", text)

    def test_the_checksum_column_can_be_left_out(self):
        spec = rs.with_on(rs.default_spec(), "figures", False)
        self.assertIn("SHA-256", "\n".join(self.build(spec)))
        spec = rs.with_option(spec, "sha", "none")
        text = "\n".join(self.build(spec))
        self.assertNotIn("SHA-256", text)
        self.assertNotIn("abababab", text)

    def test_the_calibration_is_not_said_twice_when_methods_state_it(self):
        self.details["methods"] = "Text. C 1s at 284.8 eV."
        text = "\n".join(self.build(rs.with_on(rs.default_spec(), "figures",
                                               False)))
        self.assertEqual(text.count("C 1s at 284.8 eV."), 1)
        both = "\n".join(self.build(rs.with_on(
            rs.with_on(rs.default_spec(), "figures", False), "methods", False)))
        self.assertEqual(both.count("C 1s at 284.8 eV."), 1)   # own section

    def test_nothing_chosen_is_an_error(self):
        import report
        with self.assertRaises(report.ReportError):
            self.build(rs.with_all(rs.default_spec(), False))

    def test_the_old_sections_argument_still_works(self):
        import report
        path = os.path.join(self.dir, "old.pdf")
        n = report.build_report(path, self.details, "", self.rows, self.docs,
                                self.figs, self.render, ("cover", "figures"))
        with mupdf.open(path) as d:
            text = "\n".join(p.get_text() for p in d)
        self.assertEqual(n, 4)                     # cover + three figures
        self.assertIn("Everything worked.", text)
        self.assertNotIn("Acquisition metadata", text)


@unittest.skipUnless(HAVE_PPTX, "python-pptx / matplotlib not installed")
class TestDeckFollowsTheSpec(Tmp):
    def setUp(self):
        super().setUp()
        from test_metasummary import Doc, region as mregion
        self.docs = [Doc([mregion("Survey", 160, step=1.0)], "a.vgd"),
                     Doc([mregion("Mo 3d", 40)], "b.vgd")]
        self.figs = [{"name": "One", "caption": "", "state": {}},
                     {"name": "Two", "caption": "", "state": {}}]
        self.rows = [{"name": "a.vgd", "format": "Thermo", "regions": 1,
                      "size": 2048, "sha256": "ab" * 32}]
        self.details = {"title": "The Study", "summary": "Everything worked.",
                        "methods": "Spectra were recorded.",
                        "calibration": "C 1s at 284.8 eV."}

    def png(self, number, fig):
        from matplotlib.figure import Figure
        f = Figure(figsize=(12.1, 4.95))
        f.text(0.5, 0.5, f"fig {number}")
        buf = io.BytesIO()
        f.savefig(buf, format="png", dpi=50)
        return [buf.getvalue()]

    def build(self, spec):
        import pptx_export
        from pptx import Presentation
        path = os.path.join(self.dir, "d.pptx")
        pptx_export.build_deck(path, self.details, "", self.rows, self.docs,
                               self.figs, self.png, spec=spec)
        return Presentation(path)

    @staticmethod
    def titles(prs):
        return [s.shapes.title.text if s.shapes.title is not None else ""
                for s in prs.slides]

    def test_the_order_is_the_specs_order(self):
        titles = self.titles(self.build(rs.moved(rs.default_spec(), "figures",
                                                 -5)))
        self.assertTrue(titles[0].startswith("Figure 1"))
        self.assertTrue(titles[1].startswith("Figure 2"))
        self.assertEqual(titles[-1], "Data files")

    def test_a_skipped_figure_and_a_skipped_file(self):
        spec = rs.with_child(rs.with_child(rs.default_spec(), "figures",
                                           "fig1", False),
                             "metadata", "b.vgd", False)
        titles = self.titles(self.build(spec))
        self.assertEqual([t for t in titles if t.startswith("Figure")],
                         ["Figure 1 – Two"])              # renumbered
        self.assertTrue(any("a.vgd" in t for t in titles))
        self.assertFalse(any("b.vgd" in t for t in titles))

    def test_the_calibration_rides_on_the_summary_or_stands_alone(self):
        prs = self.build(rs.default_spec())
        summary = next(s for s in prs.slides
                       if s.shapes.title.text == "Summary")
        text = " ".join(sh.text_frame.text for sh in summary.shapes
                        if sh.has_text_frame)
        self.assertIn("Energy calibration: C 1s at 284.8 eV.", text)
        self.assertNotIn("Energy calibration", self.titles(prs))
        prs = self.build(rs.with_on(rs.default_spec(), "summary", False))
        self.assertIn("Energy calibration", self.titles(prs))
        prs = self.build(rs.with_on(rs.with_on(
            rs.default_spec(), "summary", False), "calibration", False))
        self.assertNotIn("Energy calibration", self.titles(prs))

    def test_the_checksum_column(self):
        def files_text(spec):
            prs = self.build(spec)
            slide = next(s for s in prs.slides
                         if s.shapes.title.text == "Data files")
            return " ".join(c.text for sh in slide.shapes
                            if getattr(sh, "has_table", False) and sh.has_table
                            for r in sh.table.rows for c in r.cells)
        self.assertIn("SHA-256", files_text(rs.default_spec()))
        self.assertNotIn("SHA-256", files_text(rs.with_option(
            rs.default_spec(), "sha", "none")))


class TestWorkbookMember(Tmp):
    def test_the_spec_is_saved_and_an_older_workbook_has_none(self):
        import workbook as wbk
        src = os.path.join(self.dir, "a.txt")
        with open(src, "wb") as fh:
            fh.write(b"data")
        spec = rs.moved(rs.default_spec(), "files", -3)
        book = wbk.Workbook(
            files=[wbk.FileEntry(id="f1", name="a.txt", path=src)],
            report=spec)
        path = os.path.join(self.dir, "w" + wbk.EXT)
        wbk.save(path, book, None)
        self.assertEqual(rs.sanitise(wbk.load(path, os.path.join(
            self.dir, "x")).report), spec)
        book.report = {}
        wbk.save(path, book, None)
        self.assertEqual(wbk.load(path, os.path.join(self.dir, "y")).report, {})
        import zipfile
        with zipfile.ZipFile(path) as zf:
            self.assertNotIn("report.json", zf.namelist())   # optional member


try:
    import tkinter as tk
    import spectradeck as ee
    from readers.base import SpectrumFile, Region
    HAVE_APP = ee.HAVE_MPL and HAVE_PDF and HAVE_PPTX
except Exception:                                    # pragma: no cover
    HAVE_APP = False


@unittest.skipUnless(HAVE_APP, "the app, matplotlib, reportlab or pptx missing")
class TestInTheApp(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # a Workspace applies its theme to matplotlib's global settings;
        # give them back so other tests see the defaults
        import matplotlib
        cls._rc = matplotlib.rcParams.copy()
        try:
            cls.root = tk.Tk()
        except tk.TclError:
            raise unittest.SkipTest("no display")
        cls.root.withdraw()
        cls.dir = tempfile.mkdtemp()
        cls.path = os.path.join(cls.dir, "a.vgd")
        with open(cls.path, "wb") as fh:
            fh.write(b"stand-in")
        e = [280 + 0.5 * i for i in range(21)]
        r = Region(name="C 1s", index=0, offset=0, energy=e,
                   counts=[1.0] * 21, decodable=True, sample="S1",
                   photon_energy=1486.6, pass_energy=20.0, dwell=0.1,
                   step=0.5)
        f = SpectrumFile()
        f.path, f.format_name, f.regions = cls.path, "Test", [r]
        f.instrument = {"Instrument": "Test Spec"}
        f._finish()
        cls.doc, cls._load = f, ee.load_file
        ee.load_file = lambda p: cls.doc
        cls.ws = ee.Workspace(cls.root)
        cls.ws._add_file(cls.path)
        cls._boxes = (ee.messagebox.showinfo, ee.messagebox.showerror,
                      ee.messagebox.askyesno, ee.filedialog.asksaveasfilename)

    @classmethod
    def tearDownClass(cls):
        (ee.messagebox.showinfo, ee.messagebox.showerror,
         ee.messagebox.askyesno, ee.filedialog.asksaveasfilename) = cls._boxes
        ee.load_file = cls._load
        cls.root.destroy()
        import matplotlib
        matplotlib.rcParams.update(cls._rc)
        shutil.rmtree(cls.dir, ignore_errors=True)

    def setUp(self):
        ws = self.ws
        ws.details.update({"title": "App Study", "summary": "It worked."})
        ws.figures = []
        ws.set_report_spec(rs.default_spec())
        self.shown = []
        ee.messagebox.showinfo = lambda *a, **k: self.shown.append(a)
        ee.messagebox.showerror = lambda *a, **k: self.shown.append(a)
        ee.messagebox.askyesno = lambda *a, **k: False

    def test_the_inventory_reflects_the_workbook(self):
        inv = self.ws.report_inventory()
        self.assertIn("summary", inv.present_ids())
        self.assertIn("files", inv.present_ids())
        self.assertNotIn("figures", inv.present_ids())
        self.assertEqual(inv.children["metadata"][0][1], "a.vgd")

    def test_the_choice_is_remembered_and_marks_the_workbook_changed(self):
        ws = self.ws
        before = ws._signature()
        ws.set_report_spec(rs.with_on(ws.report_spec, "files", False))
        self.assertNotEqual(ws._signature(), before)
        self.assertEqual(ws.cfg["report_last"], ws.report_spec)
        ws.set_report_spec(ws.report_spec)                 # unchanged: no-op
        self.assertFalse(rs.is_on(ws.report_spec, "files"))

    def test_presets_are_kept_in_the_config(self):
        ws = self.ws
        ws.set_report_presets({"Mine": rs.BUILTIN_PRESETS["Quick look"],
                               "Everything": rs.default_spec()})
        self.assertEqual(list(ws.report_presets), ["Mine"])
        self.assertEqual(list(ws.cfg["report_presets"]), ["Mine"])
        ws.set_report_presets({})

    def dialog(self):
        import reportgen_ui
        dlg = reportgen_ui.ReportGeneratorDialog(self.root, self.ws)
        self.addCleanup(dlg.close)
        return dlg

    def test_the_dialog_lists_sections_and_children(self):
        dlg = self.dialog()
        names = [dlg.tree.set(i, "name").strip() for i in dlg.tree.get_children()]
        self.assertEqual(names[0], "Cover page")
        self.assertIn("Data files", names)
        self.assertEqual(len(dlg.rows), len(rs.SECTION_IDS))   # 1 file: no kids

    def test_preview_button_is_disabled_for_a_pptx_only_output(self):
        dlg = self.dialog()
        self.assertEqual(str(dlg.preview_btn["state"]), "normal")
        dlg.out_pdf.set(False)
        dlg.out_pptx.set(True)
        dlg._sync_preview_button()
        self.assertEqual(str(dlg.preview_btn["state"]), "disabled")
        self.assertNotEqual(dlg.preview_btn["text"], "Preview PDF")
        self.assertEqual(dlg._chosen_kinds(), {"pptx"})
        dlg.out_pdf.set(True)
        dlg._sync_preview_button()
        self.assertEqual(str(dlg.preview_btn["state"]), "normal")
        self.assertEqual(dlg.preview_btn["text"], "Preview PDF")
        self.assertEqual(dlg._chosen_kinds(), {"pdf", "pptx"})

    def test_generate_with_nothing_ticked_asks_for_a_format_first(self):
        dlg = self.dialog()
        dlg.out_pdf.set(False)
        ee.filedialog.asksaveasfilename = lambda **k: self.fail("asked to save")
        dlg.generate()
        self.assertTrue(self.shown)
        self.assertIn("Tick at least one format", self.shown[-1][1])

    def test_ticking_moving_and_selecting_change_the_apps_choice(self):
        dlg = self.dialog()
        self.assertTrue(dlg.toggle_section("summary"))
        self.assertFalse(rs.is_on(self.ws.report_spec, "summary"))
        dlg.move("files", -11)
        self.assertEqual(rs.order(self.ws.report_spec)[0], "files")
        dlg.select_all(False)
        self.assertEqual(rs.active(self.ws.report_spec), [])
        dlg.select_all(True)
        on = {i for i, _k in rs.active(self.ws.report_spec)}
        self.assertEqual(on, self.ws.report_inventory().present_ids())

    def test_an_empty_section_cannot_be_ticked(self):
        dlg = self.dialog()
        before = self.ws.report_spec
        self.assertFalse(dlg.toggle_section("figures"))     # no saved figures
        self.assertEqual(self.ws.report_spec, before)        # nothing changed

    def test_presets_apply_save_and_delete(self):
        dlg = self.dialog()
        dlg.apply_preset("Quick look")
        self.assertEqual([i for i, _k in rs.active(self.ws.report_spec)],
                         ["cover", "figures"])
        self.assertFalse(dlg.save_preset("Quick look"))    # a built-in name
        self.assertFalse(dlg.save_preset("  "))
        dlg.apply_preset("Everything")
        dlg.toggle_section("files")
        self.assertTrue(dlg.save_preset("No files"))
        self.assertIn("No files", self.ws.report_presets)
        dlg.apply_preset("Everything")
        self.assertTrue(rs.is_on(self.ws.report_spec, "files"))
        dlg.apply_preset("No files")
        self.assertFalse(rs.is_on(self.ws.report_spec, "files"))
        self.assertTrue(dlg.delete_preset("No files"))
        self.assertFalse(dlg.delete_preset("No files"))
        self.assertEqual(self.ws.report_presets, {})

    def test_the_checksum_option(self):
        dlg = self.dialog()
        dlg.set_option("sha", "none")
        self.assertEqual(rs.option(self.ws.report_spec, "sha"), "none")

    def test_generate_writes_the_pdf_the_deck_or_both(self):
        ws = self.ws
        out = os.path.join(self.dir, "out")
        os.makedirs(out, exist_ok=True)
        ws.set_report_spec(rs.with_on(rs.default_spec(), "metadata", False))
        pdf = os.path.join(out, "r.pdf")
        ee.filedialog.asksaveasfilename = lambda **k: pdf
        ws.generate_report("pdf")
        with mupdf.open(pdf) as d:
            text = "\n".join(p.get_text() for p in d)
        self.assertIn("App Study", text)
        self.assertIn("It worked.", text)
        self.assertNotIn("Acquisition metadata", text)     # switched off
        ws.generate_report("both")
        self.assertTrue(os.path.isfile(os.path.join(out, "r.pptx")))
        deck = os.path.join(out, "d.pptx")
        ee.filedialog.asksaveasfilename = lambda **k: deck
        ws.generate_report("pptx")
        self.assertTrue(os.path.isfile(deck))
        self.assertEqual(self.shown, [])                   # no errors

    @unittest.skipUnless(HAVE_DOCX, "python-docx not installed")
    def test_generate_writes_any_combination_of_formats(self):
        ws = self.ws
        out = os.path.join(self.dir, "combo")
        os.makedirs(out, exist_ok=True)
        pdf = os.path.join(out, "c.pdf")
        ee.filedialog.asksaveasfilename = lambda **k: pdf
        ws.generate_report({"pdf", "docx"})
        self.assertTrue(os.path.isfile(pdf))
        self.assertTrue(os.path.isfile(os.path.join(out, "c.docx")))
        self.assertFalse(os.path.isfile(os.path.join(out, "c.pptx")))
        self.assertEqual(self.shown, [])                   # no errors

    def test_generate_report_ignores_an_empty_kind_set(self):
        ws = self.ws
        ee.filedialog.asksaveasfilename = lambda **k: self.fail("asked to save")
        ws.generate_report(set())
        self.assertEqual(self.shown, [])

    def test_progress_total_counts_figures_and_multiplies_per_format(self):
        ws = self.ws
        ws.figures = [{"name": "F1", "caption": "", "state": ws.capture_state()}]
        spec = rs.default_spec()
        self.assertEqual(ws._report_progress_total(spec, "pdf"), 1)
        self.assertEqual(ws._report_progress_total(spec, "both"), 2)
        self.assertEqual(
            ws._report_progress_total(spec, {"pdf", "pptx", "docx"}), 3)
        ws.figures = []

    def test_cancelling_mid_generation_saves_nothing_and_says_so(self):
        ws = self.ws
        ws.checked = {id(r) for r in ws.docs[0].regions}
        ws.figures = [{"name": "F1", "caption": "", "state": ws.capture_state()}]
        out = os.path.join(self.dir, "cancel.pdf")
        ee.filedialog.asksaveasfilename = lambda **k: out
        orig_tick = ws._report_tick

        def cancel_on_first_tick(label):
            ws._gen_prog.cancelled = True
            orig_tick(label)
        ws._report_tick = cancel_on_first_tick
        try:
            ws.generate_report("pdf")
        finally:
            ws._report_tick = orig_tick
            ws.figures = []
            ws.checked = set()
        self.assertFalse(os.path.isfile(out))
        self.assertTrue(self.shown)
        self.assertEqual(self.shown[-1][0], "Report cancelled")
        self.assertIn("Nothing was saved", self.shown[-1][1])

    def test_nothing_selected_says_so_and_writes_nothing(self):
        ws = self.ws
        ws.set_report_spec(rs.with_all(rs.default_spec(), False))
        ee.filedialog.asksaveasfilename = lambda **k: self.fail("asked to save")
        ws.generate_report("pdf")
        self.assertTrue(self.shown)
        self.assertIn("Nothing selected", self.shown[0][1])

    def test_the_hand_over_report_uses_the_choice_too(self):
        ws = self.ws
        ws.set_report_spec(rs.with_on(rs.default_spec(), "summary", False))
        tmp = os.path.join(self.dir, "h.pdf")
        ws._build_report(tmp)
        with mupdf.open(tmp) as d:
            text = "\n".join(p.get_text() for p in d)
        self.assertNotIn("It worked.", text)

    def test_the_quantification_is_read_once_and_again_after_a_change(self):
        ws = self.ws
        first = ws._results()
        self.assertFalse(first)                       # no fits in this file
        self.assertIs(ws._results(), first)           # remembered
        self.assertNotIn("results", ws.report_inventory().present_ids())
        ws._ann_changed(relabel=False)
        self.assertIsNot(ws._results(), first)        # an edit reads it again

    def test_all_or_none_of_a_sections_items(self):
        ws = self.ws
        ws._has_image_pages = lambda: True
        ws._image_items = lambda: [(f"cam:a/{i}", f"Picture {i}")
                                   for i in range(8)]
        try:
            dlg = self.dialog()
            self.assertTrue(dlg.set_children("images", False))
            self.assertEqual(len(rs.skipped(ws.report_spec, "images")), 8)
            self.assertTrue(dlg.set_children("images", True))
            self.assertEqual(rs.skipped(ws.report_spec, "images"), set())
            dlg.toggle_child("images", "cam:a/3")               # pick one ...
            self.assertEqual(rs.skipped(ws.report_spec, "images"),
                             {"cam:a/3"})
            self.assertFalse(dlg.set_children("figures", False))  # no items
            self.assertFalse(dlg.set_children(None, False))
            self.assertEqual(rs.with_children(
                rs.default_spec(), "images", ["a", "b"], False)["skip"],
                {"images": ["a", "b"]})
            self.assertEqual(rs.with_children(
                rs.with_children(rs.default_spec(), "images", ["a", "b"],
                                 False), "images", ["a"], True)["skip"],
                {"images": ["b"]})
        finally:
            for name in ("_has_image_pages", "_image_items"):
                del ws.__dict__[name]

    def test_the_pdf_gets_the_title_of_every_picture_page(self):
        """What ``_report_image_pages`` hands ``build_report`` is the title of
        each page it wrote, so the contents can list each one."""
        from matplotlib.backends.backend_pdf import PdfPages
        import imagepages as ip
        import test_imagepages
        if not test_imagepages.HAVE:
            self.skipTest("matplotlib, numpy and Pillow needed")
        ws = self.ws
        ws._image_page_plan = lambda skip=(), mosaics=False: ip.plan(
            [test_imagepages.site_doc(2)])
        ws._pdf_figure_size = (11.7, 8.3)
        try:
            with PdfPages(os.path.join(self.dir, "p.pdf")) as pdf:
                titles = ws._report_image_pages(pdf)
        finally:
            del ws.__dict__["_image_page_plan"]
        self.assertEqual(titles, ["Camera pictures", "SnapMap – S1"])

    def test_the_pictures_left_out_reach_the_report_and_the_deck(self):
        ws = self.ws
        seen = []
        ws._has_image_pages = lambda: True
        ws._image_items = lambda: [("cam:a/x", "x"), ("map:a/S1", "Map")]
        ws._report_image_pages = (
            lambda pdf, skip=(), mosaics=False:
            seen.append(("pdf", set(skip), mosaics)) or 0)
        ws._deck_image_pages = (
            lambda skip=(), mosaics=False:
            seen.append(("deck", set(skip), mosaics)) or [])
        try:
            self.assertEqual(ws.report_inventory().children["images"],
                             [("cam:a/x", "x"), ("map:a/S1", "Map")])
            spec = rs.with_child(rs.default_spec(), "images", "cam:a/x",
                                 False)
            ws._build_report(os.path.join(self.dir, "i.pdf"), spec)
            if HAVE_PPTX:
                ws._build_deck(os.path.join(self.dir, "i.pptx"), spec)
        finally:
            for name in ("_has_image_pages", "_image_items",
                         "_report_image_pages", "_deck_image_pages"):
                del ws.__dict__[name]
        self.assertEqual(seen[0], ("pdf", {"cam:a/x"}, True))    # mosaics: on
        if HAVE_PPTX:
            self.assertEqual(seen[1], ("deck", {"cam:a/x"}, True))

    def test_fits_in_the_files_reach_the_report_and_the_generator(self):
        import resultspages
        from test_results import sample, three_element_level
        ws = self.ws
        made = resultspages.Results(samples=[
            sample("Film", [three_element_level(None)])])
        real, resultspages.collect = resultspages.collect, \
            lambda *a, **k: made
        self.addCleanup(setattr, resultspages, "collect", real)
        ws._ann_changed(relabel=False)
        self.assertIn("results", ws.report_inventory().present_ids())
        tmp = os.path.join(self.dir, "q.pdf")
        ws.set_report_spec(rs.default_spec())
        ws._build_report(tmp)
        with mupdf.open(tmp) as d:
            text = "\n".join(p.get_text() for p in d)
        self.assertIn("Quantification", text)
        self.assertIn("Ti 2p", text)
        ws._ann_changed(relabel=False)


if __name__ == "__main__":
    unittest.main()
