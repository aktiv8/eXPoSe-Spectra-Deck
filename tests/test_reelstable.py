"""The REELS π / (σ+π) table (reelstable.py) and where it goes: the "REELS
analysis" report section (reportspec.py; PDF, slides, Word), the hand-over CSV
(handover.py) and *Export REELS ratios* (the workspace).

Run:  python -m unittest discover tests
"""

import csv
import io
import os
import shutil
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
sys.path.insert(0, HERE)

import handover  # noqa: E402
import reels  # noqa: E402
import reelstable  # noqa: E402
import reportspec as rs  # noqa: E402
from test_glance import (Built, HAVE_DOCX, HAVE_PDF, HAVE_PPTX,  # noqa: E402
                         mupdf)
from test_iss_reels import ratio_spectrum  # noqa: E402


def stored(pi_area=300.0, sp_area=2000.0, **over):
    """A stored ratio result (what ``Annotations.pi_ratio`` holds)."""
    _ke, y, loss = ratio_spectrum(pi_area, sp_area)
    res = {**reels.pi_ratio(loss, y), "elastic": 1000.0}
    res.update(over)
    return res


def item(sample="S1", spectrum="REELS", res=None, unit="counts/s", file="a.vms"):
    return {"file": file, "sample": sample, "spectrum": spectrum, "unit": unit,
            "res": res if res is not None else stored()}


HOPG = {"hopg": {"ratio": 0.16, "label": "HOPG", "kind": "spectrum",
                 "source": "HOPG 1 keV", "pi_window": [3.5, 8.0],
                 "sp_window": [10.0, 35.0], "avg": 3, "e0": 1000.0}}


class TestTable(unittest.TestCase):
    def test_an_empty_table_is_falsy_and_has_no_text(self):
        t = reelstable.collect([], {})
        self.assertFalse(t)
        self.assertEqual(len(t), 0)
        self.assertFalse(reelstable.collect([{"file": "a", "res": None}]))
        self.assertFalse(reelstable.collect(None))

    def test_the_rows_carry_what_was_measured(self):
        t = reelstable.collect([item("A"), item("B", res=stored(0.0))], HOPG)
        self.assertEqual(len(t), 2)
        a, b = t.rows
        self.assertEqual((a.sample, a.spectrum, a.file), ("A", "REELS", "a.vms"))
        self.assertAlmostEqual(a.ratio, 0.15, delta=0.01)
        self.assertAlmostEqual(a.fraction, a.ratio / 0.16, 9)
        self.assertAlmostEqual(a.pos_pi, 5.8, delta=0.3)
        self.assertAlmostEqual(a.pos_sp, 26.0, delta=1.0)
        self.assertEqual(a.pi_window, [3.5, 8.0])
        self.assertAlmostEqual(b.ratio, 0.0, delta=0.005)
        self.assertEqual(t.unit, "counts/s")
        self.assertEqual(len(t.settings), 1)            # one setting for both

    def test_mixed_units_and_settings(self):
        other = stored(pi_window=[3.0, 8.0])
        t = reelstable.collect([item("A"), item("B", res=other, unit="counts")])
        self.assertEqual(t.unit, "")
        self.assertEqual(len(t.settings), 2)
        self.assertIn("intensity unit times eV", reelstable.units_text(t))
        one = reelstable.collect([item()])
        self.assertIn("counts/s·eV", reelstable.units_text(one))

    def test_a_missing_ratio_and_a_missing_fraction_are_blank_not_zero(self):
        t = reelstable.collect([item(res=stored(ratio=None, i_sp=-3.0))], HOPG)
        row = t.rows[0]
        self.assertIsNone(row.ratio)
        self.assertIsNone(row.fraction)
        cells = reelstable.report_rows(t)[0]
        self.assertEqual(cells[7], reelstable.DASH)
        self.assertEqual(cells[8], reelstable.DASH)
        csv_row = reelstable.csv_rows(t)[1]
        self.assertEqual(csv_row[8], "")                 # ratio
        self.assertEqual(csv_row[9], "")                 # sp2 fraction

    def test_a_number_outside_the_scale_is_flagged_never_clipped(self):
        t = reelstable.collect([item(res=stored(ratio=0.40))], HOPG)
        self.assertAlmostEqual(t.rows[0].fraction, 2.5, 9)
        self.assertTrue(any("outside" in n for n in t.rows[0].notes))
        self.assertEqual(reelstable.report_rows(t)[0][8], "2.50")
        self.assertIn("†", reelstable.report_rows(t)[0][0])

    def test_a_value_that_rounds_to_zero_has_no_minus_sign_but_the_csv_keeps_it(self):
        t = reelstable.collect([item(res=stored(i_pi=-0.004, ratio=-0.00001,
                                                i_sp=1000.0))], HOPG)
        cells = reelstable.report_rows(t)[0]
        self.assertEqual((cells[5], cells[7]), ("0.0", "0.0000"))
        csv_row = reelstable.csv_rows(t)[1]
        self.assertEqual(csv_row[6], "-0.004")
        self.assertEqual(csv_row[8], "-1e-05")
        self.assertEqual(reelstable.report_rows(reelstable.collect(
            [item(res=stored(i_pi=-85.1, ratio=-0.005))], HOPG))[0][5],
            "-85.1")

    def test_a_clean_spectrum_has_no_mark(self):
        t = reelstable.collect([item()], HOPG)
        self.assertEqual(t.rows[0].flags, [])
        self.assertNotIn("†", reelstable.report_rows(t)[0][0])
        self.assertFalse(any(p.startswith("†")
                             for p in reelstable.paragraphs(t)))

    def test_the_references_in_words(self):
        t = reelstable.collect([item()], HOPG)
        self.assertIn("HOPG 0.1600 (measured, HOPG 1 keV) = 1",
                      reelstable.reference_text(t))
        both = {**HOPG, "diamond": {"ratio": 0.01, "label": "diamond",
                                    "kind": "spectrum", "source": "",
                                    "pi_window": [3.5, 8.0],
                                    "sp_window": [10.0, 35.0], "avg": 3,
                                    "e0": 1000.0}}
        text = reelstable.reference_text(reelstable.collect([item()], both))
        self.assertIn("diamond 0.0100 (measured) = 0", text)
        typed = {"hopg": {**HOPG["hopg"], "kind": "entered",
                          "source": "lab note"}}
        self.assertIn("typed, lab note",
                      reelstable.reference_text(reelstable.collect([item()],
                                                                   typed)))
        none = reelstable.reference_text(reelstable.collect([item()], {}))
        self.assertIn("No HOPG reference", none)

    def test_no_reference_means_no_fraction_and_the_text_says_so(self):
        t = reelstable.collect([item()], {})
        self.assertIsNone(t.rows[0].fraction)
        self.assertEqual(reelstable.report_rows(t)[0][8], reelstable.DASH)

    def test_a_reference_measured_other_wise_is_said_per_spectrum(self):
        t = reelstable.collect([item(res=stored(pi_window=[3.0, 8.0]))], HOPG)
        self.assertIsNone(t.rows[0].fraction)
        self.assertIn("windows", t.rows[0].problem)
        lines = reelstable.issue_lines(t)
        self.assertTrue(any("no sp² fraction" in l for l in lines))

    def test_problems_of_one_kind_are_one_line_naming_the_spectra(self):
        low = [item(f"S{i}", res=stored(0.0, 2000.0, dip_pi=0.2 + i / 10))
               for i in range(1, 4)]
        t = reelstable.collect(low, HOPG)
        lines = reelstable.issue_lines(t)
        dips = [l for l in lines if "dips" in l]
        self.assertEqual(len(dips), 1)                   # not three lines
        self.assertIn("all 3 spectra", dips[0])
        one = reelstable.issue_lines(reelstable.collect(
            [item("A"), item("B", res=stored(dip_pi=0.5))], HOPG))
        self.assertTrue(any(l.endswith("— B / REELS") for l in one), one)

    def test_a_percentage_that_varies_is_given_as_its_range(self):
        rows = [item(f"S{i}", res=stored(0.0, 2000.0, dip_pi=d))
                for i, d in enumerate((0.22, 0.78, 3.78))]
        lines = reelstable.issue_lines(reelstable.collect(rows, HOPG))
        dips = [l for l in lines if "π curve dips" in l]
        self.assertEqual(len(dips), 1)
        self.assertIn("dips 22–378 %", dips[0])
        same = reelstable.issue_lines(reelstable.collect(
            [item("A", res=stored(dip_pi=0.5)),
             item("B", res=stored(dip_pi=0.5))], HOPG))
        self.assertTrue(any("dips 50 %" in l for l in same), same)

    def test_the_beam_energy_note_does_not_print_one_spectrum_energies(self):
        rows = [item("A", res=stored(elastic=2000.0)),
                item("B", res=stored(elastic=3000.0))]
        lines = reelstable.issue_lines(reelstable.collect(rows, HOPG))
        beam = [l for l in lines if "beam energy" in l]
        self.assertEqual(len(beam), 1)
        self.assertNotIn("2000 eV", beam[0])
        self.assertNotIn("3000 eV", beam[0])
        one = reelstable.issue_lines(reelstable.collect(rows[:1], HOPG))
        self.assertTrue(any("2000 eV" in l for l in one))   # alone: kept

    def test_the_text_that_goes_with_the_table(self):
        t = reelstable.collect([item()], HOPG)
        paras = reelstable.paragraphs(t)
        self.assertIn("3.5–8 eV", paras[0])
        self.assertIn("centred", paras[0])
        self.assertIn("Elastic is the kinetic energy", paras[1])
        self.assertEqual(paras[-1], reelstable.CAVEAT)
        self.assertIn("not an sp² percentage", paras[-1])


class TestCsv(unittest.TestCase):
    def table(self):
        return reelstable.collect(
            [item("A", "REELS Survey", file="x.experiment"),
             item("B", res=stored(0.0))], HOPG)

    def test_the_columns_follow_xpsview_with_the_additions(self):
        head = reelstable.csv_rows(self.table())[0]
        # HarwellXPS's reels_pi_sigma_ratio.csv: elastic peak, pi, sigma+pi,
        # the two areas, the ratio, the two windows
        want = ["Elastic peak KE (eV)", "pi position (eV)",
                "sigma+pi position (eV)"]
        self.assertEqual(head[3:6], want)
        self.assertTrue(head[6].startswith("I_pi ("))
        self.assertTrue(head[7].startswith("I_sigma+pi ("))
        self.assertEqual(head[8], "I_pi/I_sigma+pi")
        self.assertEqual(head[10:12], ["pi window (eV)", "sigma+pi window (eV)"])
        self.assertEqual(head[:3], ["File", "Sample", "Spectrum"])
        self.assertEqual(head[-3:], ["HOPG ratio", "Diamond ratio", "Notes"])

    def test_bytes_are_utf8_with_a_bom_and_read_back(self):
        data = reelstable.csv_bytes(self.table())
        self.assertTrue(data.startswith(b"\xef\xbb\xbf"))
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[1][:3], ["x.experiment", "A", "REELS Survey"])
        self.assertEqual(rows[1][10], "3.5-8")
        self.assertEqual(rows[1][11], "10-35")
        self.assertEqual(rows[1][12], "3")
        self.assertEqual(rows[1][13], "0.16")
        self.assertEqual(rows[1][14], "")                # no diamond stored
        self.assertAlmostEqual(float(rows[1][8]),
                               self.table().rows[0].ratio, 5)
        self.assertEqual(len({len(r) for r in rows}), 1)  # rectangular

    def test_a_note_with_a_comma_or_a_quote_survives(self):
        t = reelstable.collect([item(res=stored(dip_pi=0.9))], HOPG)
        data = reelstable.csv_bytes(t).decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(data)))
        self.assertIn("dips 90 %", rows[1][-1])

    def test_greek_letters_in_names_are_kept(self):
        t = reelstable.collect([item("Kα film", "π test")], HOPG)
        data = reelstable.csv_bytes(t).decode("utf-8-sig")
        self.assertIn("Kα film", data)
        self.assertIn("π test", data)


class TestSection(unittest.TestCase):
    def test_it_is_a_report_section_after_the_quantification(self):
        order = list(rs.SECTION_IDS)
        self.assertEqual(order.index("reels"), order.index("results") + 1)
        self.assertEqual(rs.LABELS["reels"], "REELS analysis")
        self.assertIn("reels", rs.SHORT)
        self.assertIn("reels", rs.HINTS)
        self.assertIn("reels", rs.LEGACY_ORDER)

    def test_a_new_spec_has_it_on_and_an_old_one_does_not_gain_it(self):
        self.assertTrue(rs.is_on(rs.default_spec(), "reels"))
        old = rs.default_spec()
        old["sections"] = [s for s in old["sections"] if s["id"] != "reels"]
        self.assertFalse(rs.is_on(rs.sanitise(old), "reels"))
        self.assertTrue(rs.is_on(rs.sanitise({}), "reels"))     # garbage: on
        self.assertEqual(rs.order(rs.sanitise(old)).index("reels"),
                         rs.order(rs.sanitise(old)).index("results") + 1)

    def test_presets(self):
        on = lambda name: rs.is_on(rs.all_presets({})[name], "reels")
        self.assertTrue(on("Everything"))
        self.assertTrue(on("Customer report"))
        self.assertFalse(on("Audit trail"))
        self.assertFalse(on("Quick look"))
        for name, spec in rs.all_presets({}).items():
            self.assertEqual(sorted(rs.order(spec)), sorted(rs.SECTION_IDS))

    def test_the_old_section_names_do_not_switch_it_on(self):
        for kind in ("pdf", "deck"):
            spec = rs.spec_from_sections(("cover", "title", "metadata",
                                          "images", "figures"), kind)
            self.assertFalse(rs.is_on(spec, "reels"), kind)

    def inv(self, **kw):
        args = dict(details={"summary": "s"}, methods_text="m",
                    calibration="c", file_rows=[{"name": "a"}], docs=[],
                    figures=[], has_images=False)
        args.update(kw)
        return rs.inventory(**args)

    def test_the_inventory_says_what_there_is_or_why_not(self):
        t = reelstable.collect([item("A"), item("B")], HOPG)
        inv = self.inv(reels=t)
        self.assertIn("reels", inv.present_ids())
        self.assertEqual(inv.summary("reels"), "2 spectra")
        one = self.inv(reels=reelstable.collect([item()], HOPG))
        self.assertEqual(one.summary("reels"), "1 spectrum")
        none = self.inv()
        self.assertNotIn("reels", none.present_ids())
        self.assertIn("REELS", none.summary("reels"))
        empty = self.inv(reels=reelstable.collect([], {}))
        self.assertNotIn("reels", empty.present_ids())

    def test_describe_names_it_when_it_is_on(self):
        t = reelstable.collect([item()], HOPG)
        text = rs.describe(rs.default_spec(), self.inv(reels=t))
        self.assertIn("REELS analysis", text)


class Base(Built):
    def setUp(self):
        super().setUp()
        self.table = reelstable.collect(
            [item("Film A", "REELS Survey"),
             item("Film B", res=stored(0.0, 2000.0, dip_pi=0.5))], HOPG)


@unittest.skipUnless(HAVE_PDF, "reportlab / PyMuPDF / matplotlib missing")
class TestPdf(Base):
    def build(self, spec=None, **kw):
        import report
        self.k = getattr(self, "k", 0) + 1
        path = os.path.join(self.dir, f"r{self.k}.pdf")
        kw.setdefault("glance", self.facts)
        kw.setdefault("timing", self.timing)
        kw.setdefault("reels", self.table)
        report.build_report(path, self.details, "", self.rows, self.docs, [],
                            lambda *a: 0, spec=spec or rs.default_spec(), **kw)
        self.doc = mupdf.open(path)
        self.addCleanup(self.doc.close)
        return [p.get_text() for p in self.doc]

    def test_the_table_and_its_text_are_there(self):
        text = "\n".join(self.build())
        for token in ("REELS analysis", "Film A", "Film B", "REELS Survey",
                      "Iπ/Iσ+π", "1000.00", "HOPG 0.1600",
                      "not an sp² percentage", "† the π area is negative",
                      "the π curve dips 50 %"):
            self.assertIn(token, text)

    def test_it_is_in_the_contents_and_the_bookmarks(self):
        pages = self.build()
        toc = {t: p for _l, t, p in self.doc.get_toc()}
        self.assertIn("REELS analysis", toc)
        self.assertIn("REELS analysis", pages[toc["REELS analysis"] - 1])
        self.assertIn("REELS analysis", pages[1])           # printed contents

    def test_it_comes_after_the_quantification_slot_and_before_methods(self):
        self.build()
        names = [t for _l, t, _p in self.doc.get_toc()]
        self.assertLess(names.index("REELS analysis"), names.index("Methods"))

    def test_a_tick_removes_it(self):
        pages = self.build(rs.with_on(rs.default_spec(), "reels", False))
        self.assertNotIn("REELS analysis", "\n".join(pages))

    def test_nothing_measured_leaves_it_out(self):
        for table in (None, reelstable.collect([], {})):
            self.assertNotIn("REELS analysis", "\n".join(
                self.build(reels=table)))

    def test_an_old_report_does_not_gain_the_section(self):
        old = rs.default_spec()
        old["sections"] = [s for s in old["sections"] if s["id"] != "reels"]
        self.assertNotIn("REELS analysis",
                         "\n".join(self.build(rs.sanitise(old))))


@unittest.skipUnless(HAVE_PPTX, "python-pptx / PyMuPDF not installed")
class TestDeck(Base):
    def build(self, spec=None, **kw):
        import pptx_export
        from pptx import Presentation
        self.k = getattr(self, "k", 0) + 1
        path = os.path.join(self.dir, f"d{self.k}.pptx")
        kw.setdefault("glance", self.facts)
        kw.setdefault("timing", self.timing)
        kw.setdefault("reels", self.table)
        pptx_export.build_deck(path, self.details, "", self.rows, self.docs,
                               [], lambda n, f: [],
                               spec=spec or rs.default_spec(), **kw)
        return list(Presentation(path).slides)

    @staticmethod
    def title(slide):
        return slide.shapes.title.text if slide.shapes.title is not None \
            else ""

    @staticmethod
    def table_cells(slide):
        return [[c.text for c in row.cells] for sh in slide.shapes
                if sh.has_table for row in sh.table.rows]

    def test_a_table_slide_and_a_text_slide(self):
        slides = self.build()
        titles = [self.title(s) for s in slides]
        i = titles.index("REELS π / (σ+π) intensity ratio")
        cells = self.table_cells(slides[i])
        self.assertEqual(cells[0], reelstable.header(self.table))
        self.assertEqual(cells[1][0], "Film A")
        self.assertTrue(cells[2][0].startswith("Film B"))
        self.assertEqual(len(cells), 3)
        words = " ".join(
            sh.text_frame.text for sl, t in zip(slides, titles)
            if t.startswith("REELS π / (σ+π): what the numbers are")
            for sh in sl.shapes if sh.has_text_frame)
        self.assertIn("not an sp² percentage", words)
        self.assertIn("HOPG 0.1600", words)

    def test_it_is_in_the_contents(self):
        slides = self.build()
        contents = next(s for s in slides if self.title(s) == "Contents")
        names = [r[0].strip() for r in self.table_cells(contents)]
        self.assertTrue(any("REELS" in n for n in names), names)

    def test_a_long_table_continues_on_another_slide(self):
        many = reelstable.collect([item(f"S{i:02d}") for i in range(40)], HOPG)
        slides = self.build(reels=many)
        titles = [self.title(s) for s in slides]
        self.assertIn("REELS π / (σ+π) intensity ratio (continued)", titles)

    def test_a_tick_removes_it(self):
        titles = [self.title(s) for s in self.build(
            rs.with_on(rs.default_spec(), "reels", False))]
        self.assertFalse([t for t in titles if "REELS" in t])

    def test_nothing_measured_leaves_it_out(self):
        titles = [self.title(s) for s in self.build(reels=None)]
        self.assertFalse([t for t in titles if "REELS" in t])


@unittest.skipUnless(HAVE_DOCX, "python-docx not installed")
class TestDocx(Base):
    def build(self, spec=None, **kw):
        import docx_export
        from docx import Document
        self.k = getattr(self, "k", 0) + 1
        path = os.path.join(self.dir, f"d{self.k}.docx")
        kw.setdefault("glance", self.facts)
        kw.setdefault("timing", self.timing)
        kw.setdefault("reels", self.table)
        docx_export.build_document(
            path, self.details, "", self.rows, self.docs, [],
            lambda n, f: [], spec=spec or rs.default_spec(), **kw)
        return Document(path)

    @staticmethod
    def headings(doc):
        return [p.text for p in doc.paragraphs
                if p.style.name == "Title" or p.style.name.startswith(
                    "Heading")]

    def test_a_heading_a_table_and_the_text(self):
        doc = self.build()
        heads = self.headings(doc)
        self.assertIn("REELS analysis", heads)
        self.assertIn("π / (σ+π) intensity ratio", heads)
        self.assertLess(heads.index("REELS analysis"), heads.index("Methods"))
        tables = [[[c.text for c in row.cells] for row in t.rows]
                  for t in doc.tables]
        mine = next(t for t in tables
                    if t[0] == reelstable.header(self.table))
        self.assertEqual(mine[1][0], "Film A")
        self.assertTrue(mine[2][0].startswith("Film B"))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("not an sp² percentage", text)
        self.assertIn("HOPG 0.1600", text)

    def test_a_tick_removes_it_and_nothing_measured_leaves_it_out(self):
        self.assertNotIn("REELS analysis", self.headings(
            self.build(rs.with_on(rs.default_spec(), "reels", False))))
        self.assertNotIn("REELS analysis", self.headings(
            self.build(reels=None)))


class TestHandover(unittest.TestCase):
    def test_the_part_is_the_csv(self):
        t = reelstable.collect([item("A"), item("B")], HOPG)
        parts = handover.reels_parts(t)
        self.assertEqual([p.arc for p in parts], [handover.REELS_ARC])
        self.assertEqual(parts[0].data, reelstable.csv_bytes(t))
        self.assertTrue(parts[0].data.startswith(b"\xef\xbb\xbf"))

    def test_nothing_measured_gives_no_part(self):
        self.assertEqual(handover.reels_parts(reelstable.collect([], {})), [])
        self.assertEqual(handover.reels_parts(None), [])

    def test_the_readme_explains_it_only_when_it_is_there(self):
        t = reelstable.collect([item()], HOPG)
        with_it = handover.readme_text({}, handover.reels_parts(t))
        self.assertIn("About the REELS ratios", with_it)
        self.assertIn("not an sp² percentage", with_it)
        self.assertIn(handover.REELS_ARC, with_it)
        self.assertNotIn("About the REELS ratios",
                         handover.readme_text({}, []))

    def test_it_goes_into_the_zip_with_its_checksum(self):
        import zipfile
        t = reelstable.collect([item()], HOPG)
        d = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, d, True)
        path = os.path.join(d, "h.zip")
        names = handover.write_zip(path, "pkg", handover.reels_parts(t))
        self.assertIn("pkg/" + handover.REELS_ARC, names)
        with zipfile.ZipFile(path) as zf:
            sums = zf.read("pkg/SHA256SUMS.txt").decode()
            self.assertIn(handover.REELS_ARC, sums)


if __name__ == "__main__":
    unittest.main()
