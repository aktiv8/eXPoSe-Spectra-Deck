"""Camera pictures and SnapMaps as report pages and slides (imagepages.py) and
their place in the PDF report and the deck. Uses the synthetic maps and
pictures of test_htmlbrowser; needs numpy, Pillow and matplotlib."""

import copy
import io
import os
import shutil
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
sys.path.insert(0, HERE)

try:
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib.figure import Figure
    import numpy as np
    import PIL  # noqa: F401
    HAVE = True
except Exception:                                    # pragma: no cover
    HAVE = False

if HAVE:
    import imagepages as ip
    import pptx_export
    import report
    from readers import ImageBlob
    from readers.imaging import encode_png
    from test_htmlbrowser import (CAMERA_CAL, camera_blob, cube_region, doc,
                                  map_cube, region)


def two_pictures(n=2):
    """``n`` camera pictures at the map's position (times differ)."""
    return [camera_blob(f"S1 Pt #001a  1{i}:00") for i in range(n)]


def site_doc(n_pictures=1, with_map=True):
    regs = [region("C 1s", "S2")]
    if with_map:
        regs.append(cube_region(map_cube()))
    return doc("a.vgd", regs, positions={"S1": (10.0, 20.0), "S2": (10.1, 20.1)},
               images=two_pictures(n_pictures) if n_pictures else [])


@unittest.skipUnless(HAVE, "matplotlib, numpy and Pillow needed")
class TestPlan(unittest.TestCase):
    def test_sheets_come_before_map_pages(self):
        pages = ip.plan([site_doc(2)])
        self.assertEqual([(p.kind, p.title, p.n_items) for p in pages],
                         [("camera", "Camera pictures", 2),
                          ("maps", "SnapMap – S1", 1)])
        self.assertTrue(ip.available([site_doc(2)]))

    def test_long_runs_are_split_into_numbered_sheets(self):
        pages = ip.plan([site_doc(7, with_map=False)], per_sheet=3)
        self.assertEqual([(p.title, p.n_items) for p in pages],
                         [("Camera pictures (1 of 3)", 3),
                          ("Camera pictures (2 of 3)", 3),
                          ("Camera pictures (3 of 3)", 1)])

    def test_only_calibrated_pictures_and_real_maps_count(self):
        plain = ImageBlob("holder", 0, b"x", True)
        d = doc("a.vms", [region("C 1s", "S1")], images=[plain])
        self.assertEqual(ip.plan([d]), [])
        self.assertFalse(ip.available([d]))
        self.assertTrue(ip.available([site_doc(0)]))          # a map alone
        self.assertEqual([p.kind for p in ip.plan([site_doc(0)])], ["maps"])

    def test_a_map_finds_the_picture_taken_at_its_position(self):
        pages = ip.plan([site_doc(1)])
        # planned internally: the map page's notes name the camera picture
        self.assertIn("Camera picture: S1 Pt #001a", pages[1].notes())
        self.assertNotIn("Camera picture", ip.plan([site_doc(0)])[0].notes())

    def test_names_and_energies_are_shown_as_the_app_shows_them(self):
        def display(r):
            q = copy.copy(r)
            q.name = "Renamed " + r.name
            q.energy = [e + 1.5 for e in r.energy]
            return q
        pages = ip.plan([site_doc(1)], label_of=lambda p, s: "Lab " + s,
                        display=display)
        self.assertEqual(pages[1].title, "SnapMap – Lab S1")
        note = pages[1].notes()
        self.assertIn("Renamed C 1s: counts summed over", note)
        cube = map_cube()
        lo, hi = ip.window_of(display(cube_region(cube)), cube)
        import snapmap
        w = snapmap.default_window(cube)
        self.assertAlmostEqual(lo, w[0] + 1.5)
        self.assertAlmostEqual(hi, w[1] + 1.5)
        # and the picture is labelled with the shown names too
        fig = Figure(figsize=(11.7, 8.3))
        pages[0].draw(fig)
        texts = [t.get_text() for t in fig.axes[0].texts]
        self.assertIn("Lab S1", texts)                 # the map's outline
        self.assertIn("Lab S2", texts)                 # another analysis point

    def test_notes_describe_the_pictures(self):
        note = ip.plan([site_doc(2)])[0].notes()
        self.assertEqual(len(note.splitlines()), 2)
        self.assertIn("stage 10.000, 20.000 mm", note)
        self.assertIn("field of view 3.2 x 2.4 mm", note)


@unittest.skipUnless(HAVE, "matplotlib, numpy and Pillow needed")
class TestPicking(unittest.TestCase):
    """Which pictures and map sites go in: the Report generator's children."""

    def test_every_picture_and_map_site_has_a_key_and_a_label(self):
        got = ip.items([site_doc(2)])
        self.assertEqual(len(got), 3)
        keys = [k for k, _l in got]
        self.assertEqual(len(set(keys)), 3)
        self.assertEqual([k.split(":")[0] for k in keys],
                         ["cam", "cam", "map"])
        self.assertIn("SnapMap – S1", [l for _k, l in got])
        self.assertTrue(all("a.vgd" in k for k in keys))
        # the picture's name already says its sample: not said twice
        cam = next(l for k, l in got if k.startswith("cam:"))
        self.assertEqual(cam.count("S1"), 1)

    def test_uncalibrated_pictures_are_not_offered(self):
        plain = ImageBlob("holder", 0, b"x", True)
        d = doc("a.vms", [region("C 1s", "S1")], images=[plain])
        self.assertEqual(ip.items([d]), [])

    def test_a_picture_left_out_is_not_on_the_sheet(self):
        d = site_doc(2)
        first = ip.items([d])[0][0]
        pages = ip.plan([d], skip=[first])
        self.assertEqual([(p.kind, p.n_items) for p in pages],
                         [("camera", 1), ("maps", 1)])
        self.assertNotIn(first.split("/")[1], pages[0].notes())

    def test_all_pictures_left_out_leave_no_sheet_and_the_map_page_loses_its_photo(
            self):
        d = site_doc(2)
        cams = [k for k, _l in ip.items([d]) if k.startswith("cam:")]
        pages = ip.plan([d], skip=cams)
        self.assertEqual([p.kind for p in pages], ["maps"])
        self.assertNotIn("Camera picture", pages[0].notes())
        # the same map page with its picture, for comparison
        self.assertIn("Camera picture", ip.plan([d])[1].notes())

    def test_a_map_site_left_out_has_no_page(self):
        d = site_doc(1)
        site = next(k for k, _l in ip.items([d]) if k.startswith("map:"))
        self.assertEqual([p.kind for p in ip.plan([d], skip=[site])],
                         ["camera"])

    def test_the_choice_is_per_file(self):
        a, b = site_doc(1), site_doc(1)
        b.path = "b.vgd"
        keys = [k for k, _l in ip.items([a, b])]
        self.assertEqual(len(set(keys)), 4)
        pages = ip.plan([a, b], skip=[k for k in keys if "b.vgd" in k])
        self.assertEqual([(p.kind, p.n_items) for p in pages],
                         [("camera", 1), ("maps", 1)])

    def test_the_generator_lists_them_as_children_of_pictures(self):
        import reportspec as rs
        items = ip.items([site_doc(2)])
        inv = rs.inventory({}, "", "", [], [], [], True, image_items=items)
        self.assertEqual(inv.children["images"], items)
        self.assertEqual(inv.summary("images"), "3 pictures and maps")
        spec = rs.with_child(rs.default_spec(), "images", items[0][0], False)
        self.assertEqual(rs.skipped(spec, "images"), {items[0][0]})
        self.assertIn("Pictures (2 of 3)", rs.describe(spec, inv)
                      .replace("Camera pictures, SnapMaps and image maps", "Pictures"))


@unittest.skipUnless(HAVE, "matplotlib, numpy and Pillow needed")
class TestDrawing(unittest.TestCase):
    def draw(self, page, size=(11.7, 8.3), rect=(0.0, 0.05, 1.0, 0.93)):
        fig = Figure(figsize=size, dpi=50)
        page.draw(fig, rect)
        return fig

    def test_a_sheet_has_one_axes_per_picture_of_a_fixed_size(self):
        full = self.draw(ip.plan([site_doc(3, False)], per_sheet=3)[0])
        part = self.draw(ip.plan([site_doc(1, False)], per_sheet=3)[0])
        self.assertEqual(len(full.axes), 3)
        self.assertEqual(len(part.axes), 1)
        a, b = full.axes[0].get_position(), part.axes[0].get_position()
        self.assertAlmostEqual(a.width, b.width)       # same cell size

    def test_a_map_page_has_picture_map_and_colour_bar(self):
        fig = self.draw(ip.plan([site_doc(1)])[1])
        self.assertEqual(len(fig.axes), 3)             # picture, map, colour bar
        fig = self.draw(ip.plan([site_doc(0)])[0])
        self.assertEqual(len(fig.axes), 2)             # map, colour bar
        titles = [a.get_title() for a in fig.axes]
        self.assertTrue(any(t.startswith("C 1s") for t in titles))

    def test_drawing_stays_inside_the_rectangle_it_was_given(self):
        rect = (0.0, 0.05, 1.0, 0.93)
        for page in ip.plan([site_doc(2)]):
            fig = self.draw(page, rect=rect)
            for ax in fig.axes:
                pos = ax.get_position()
                self.assertGreaterEqual(pos.x0, rect[0] - 1e-9)
                self.assertLessEqual(pos.x1, rect[2] + 1e-9)
                self.assertGreaterEqual(pos.y0, rect[1] - 1e-9)
                self.assertLessEqual(pos.y1, rect[3] + 1e-9)

    def test_slide_sized_pages_draw_too(self):
        for page in ip.plan([site_doc(3)], per_sheet=3):
            fig = self.draw(page, size=pptx_export.FIGURE_SIZE,
                            rect=(0.0, 0.0, 1.0, 1.0))
            buf = io.BytesIO()
            fig.savefig(buf, format="png")
            self.assertTrue(buf.getvalue().startswith(b"\x89PNG"))

    def test_pictures_are_shrunk_for_the_page(self):
        w, h = 1600, 1200
        blob = ImageBlob("big", 0, encode_png(w, h, bytes(w * h * 3)), False,
                         fmt="png", sample="S1",
                         calib=dict(CAMERA_CAL, width=w, height=h,
                                    um_per_px_x=5.0, um_per_px_y=5.0))
        d = doc("a.vgd", [region("C 1s", "S1")], positions={"S1": (10.0, 20.0)},
                images=[blob])
        page = ip.plan([d])[0]
        fig = self.draw(page)
        img = fig.axes[0].images[0].get_array()
        self.assertEqual(img.shape[1], ip.PICTURE_PX)
        # ... but the axes still speak in the full picture's pixels
        self.assertEqual(fig.axes[0].get_xlim(), (0, w))

    def test_a_map_outline_replaces_the_marker_of_a_map_site(self):
        fig = self.draw(ip.plan([site_doc(1, True)], per_sheet=1)[0])
        ax = fig.axes[0]
        texts = [t.get_text() for t in ax.texts]
        self.assertIn("S1", texts)                     # the outline's label
        self.assertEqual(len(ax.patches) >= 1, True)
        # S1 is a map site: no second, round marker for it
        self.assertEqual(sum(1 for t in texts if t == "S1"), 1)


def imaging_doc(n_maps, positions=("Grid",)):
    """A loaded Kratos ``.kal`` of ``n_maps`` imaging maps at each position."""
    from readers import load_file
    from test_kratosmap import kal_map, kal_position, kal_text
    objs, k = [], 0
    for pos in positions:
        k += 1
        objs.append(kal_position(pos, k))
        for j in range(n_maps):
            k += 1
            objs.append(kal_map("Au 4f", k, z_m=(900 + 30 * j) / 1e6,
                                when=f"17/09/08 09:{10 + j:02d}:00"))
    tmp = tempfile.mkdtemp()
    path = os.path.join(tmp, "maps.kal")
    with open(path, "w") as fh:
        fh.write(kal_text(*objs))
    d = load_file(path)
    shutil.rmtree(tmp, True)
    return d


@unittest.skipUnless(HAVE, "matplotlib, numpy and Pillow needed")
class TestImageMaps(unittest.TestCase):
    """Kratos imaging maps: pages of single-energy images per stage position."""

    def test_one_position_is_titled_by_name_and_long_series_are_numbered(self):
        pages = ip.plan([imaging_doc(3)])
        self.assertEqual([(p.kind, p.title, p.n_items) for p in pages],
                         [("images", "Image maps – Grid", 3)])
        pages = ip.plan([imaging_doc(14)])
        self.assertEqual([(p.title, p.n_items) for p in pages],
                         [("Image maps – Grid (1 of 2)", ip.IMAGES_PER_PAGE),
                          ("Image maps – Grid (2 of 2)", 2)])

    def test_positions_share_a_page_and_a_small_one_is_not_split(self):
        pages = ip.plan([imaging_doc(3, ("A", "B"))])
        self.assertEqual([(p.title, p.n_items) for p in pages],
                         [("Image maps", 6)])
        # 7 + 7 do not fit on one page of 12: the second starts a new page
        pages = ip.plan([imaging_doc(7, ("A", "B"))])
        self.assertEqual([(p.title, p.n_items) for p in pages],
                         [("Image maps (1 of 2)", 7),
                          ("Image maps (2 of 2)", 7)])
        # each cell then names its position
        fig = Figure(figsize=(11.7, 8.3), dpi=50)
        pages[1].draw(fig)
        titles = [a.get_title() for a in fig.axes if a.images]
        self.assertTrue(titles)
        self.assertTrue(all(t.startswith("B" + chr(10)) for t in titles))

    def test_a_position_left_out_is_not_on_the_page(self):
        d = imaging_doc(2, ("A", "B"))
        a = next(k for k, l in ip.items([d]) if l.endswith("A"))
        pages = ip.plan([d], skip=[a])
        self.assertEqual([(p.title, p.n_items) for p in pages],
                         [("Image maps – B", 2)])
        self.assertEqual(ip.plan([d], skip=[k for k, _l in ip.items([d])]), [])

    def test_names_come_from_the_label_function(self):
        pages = ip.plan([imaging_doc(1)], label_of=lambda p, s: "Lab " + s)
        self.assertEqual(pages[0].title, "Image maps – Lab Grid")

    def test_each_image_has_its_own_axes_scale_bar_and_colour_bar(self):
        page = ip.plan([imaging_doc(3)])[0]
        fig = Figure(figsize=(11.7, 8.3), dpi=50)
        page.draw(fig, (0.0, 0.03, 1.0, 0.93))
        images = [a for a in fig.axes if a.images]
        self.assertEqual(len(images), 3)
        self.assertEqual(len(fig.axes), 6)             # image + colour bar each
        texts = [t.get_text() for t in images[0].texts]
        self.assertTrue(any("(approx.)" in t for t in texts))
        self.assertIn("Au 4f", images[0].get_title())
        pos = [a.get_position() for a in images]
        full = ip.plan([imaging_doc(12)])[0]
        fig2 = Figure(figsize=(11.7, 8.3), dpi=50)
        full.draw(fig2, (0.0, 0.03, 1.0, 0.93))
        # a part-filled page keeps the cell size of a full one
        self.assertAlmostEqual(pos[0].width,
                               [a for a in fig2.axes if a.images][0]
                               .get_position().width)

    def test_drawing_stays_inside_the_rectangle_and_saves_for_a_slide(self):
        import pptx_export
        rect = (0.0, 0.03, 1.0, 0.93)
        page = ip.plan([imaging_doc(12)])[0]
        fig = Figure(figsize=(11.7, 8.3), dpi=50)
        page.draw(fig, rect)
        for ax in fig.axes:
            pos = ax.get_position()
            self.assertGreaterEqual(pos.x0, rect[0] - 1e-9)
            self.assertLessEqual(pos.x1, rect[2] + 1e-9)
            self.assertGreaterEqual(pos.y0, rect[1] - 1e-9)
            self.assertLessEqual(pos.y1, rect[3] + 1e-9)
        fig = Figure(figsize=pptx_export.FIGURE_SIZE, dpi=50)
        page.draw(fig, (0.0, 0.0, 1.0, 1.0))
        buf = io.BytesIO()
        fig.savefig(buf, format="png")
        self.assertTrue(buf.getvalue().startswith(b"\x89PNG"))

    def test_notes_say_the_pixel_size_is_approximate(self):
        note = ip.plan([imaging_doc(2)])[0].notes()
        self.assertIn("Image maps (2 on this page; Grid)", note)
        self.assertIn("approximate", note)
        self.assertEqual(len(note.splitlines()), 3)    # summary + one per image

    def test_the_pdf_has_the_image_pages_with_their_titles(self):
        tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, tmp, True)

        def render_images(pdf):
            out = []
            for pg in ip.plan([imaging_doc(7, ("A", "B"))]):
                fig = Figure(figsize=(11.7, 8.3))
                pg.draw(fig, (0.0, 0.03, 1.0, 0.93))
                fig.text(0.03, 0.975, pg.title)
                pdf.savefig(fig)
                out.append(pg.title)
            return out

        path = os.path.join(tmp, "r.pdf")
        n = report.build_report(path, {"title": "Study"}, "", [], [], [],
                                lambda *a: 0, ("cover", "images"), render_images)
        self.assertEqual(n, 3)                          # cover + two image pages
        try:
            import pymupdf as mu
        except ImportError:
            import fitz as mu
        with mu.open(path) as d:
            text = [p.get_text() for p in d]
        self.assertIn("Image maps (1 of 2)", text[1])
        self.assertIn("Image maps (2 of 2)", text[2])


@unittest.skipUnless(HAVE, "matplotlib, numpy and Pillow needed")
class TestInReportAndDeck(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp, True)
        self.details = {"title": "Study"}

    def render_images(self, pdf):
        n = 0
        for pg in ip.plan([site_doc(2)]):
            fig = Figure(figsize=(11.7, 8.3))
            pg.draw(fig, (0.0, 0.03, 1.0, 0.93))
            fig.text(0.03, 0.975, pg.title)
            pdf.savefig(fig)
            n += 1
        return n

    def render_figure(self, pdf, number, figure):
        fig = Figure(figsize=(11.7, 8.3))
        fig.text(0.5, 0.5, f"FIGURE {number}")
        pdf.savefig(fig)
        return 1

    def pages(self, path):
        try:
            import pymupdf as mu
        except ImportError:
            import fitz as mu
        with mu.open(path) as d:
            return [p.get_text() for p in d]

    def build(self, sections, render_images="default"):
        path = os.path.join(self.tmp, "r.pdf")
        ri = self.render_images if render_images == "default" else render_images
        n = report.build_report(path, self.details, "", [], [],
                                [{"name": "F", "caption": "", "state": {}}],
                                self.render_figure, sections, ri)
        return n, self.pages(path)

    def test_images_sit_between_metadata_and_figures(self):
        self.assertEqual(report.SECTIONS,
                         ("cover", "metadata", "images", "figures"))
        n, text = self.build(("cover", "images", "figures"))
        self.assertEqual(n, len(text))
        self.assertEqual(n, 4)                         # cover, sheet, map, figure
        self.assertIn("Camera pictures", text[1])
        self.assertIn("SnapMap", text[2])
        self.assertIn("FIGURE 1", text[3])
        self.assertIn("report page 2 of 4", text[1])   # numbered with the rest

    def test_images_are_left_out_when_not_chosen_or_not_offered(self):
        n, text = self.build(("cover", "figures"))
        self.assertEqual(n, 2)
        n, text = self.build(("cover", "images", "figures"), render_images=None)
        self.assertEqual(n, 2)
        with self.assertRaises(report.ReportError):
            self.build(("images",), render_images=None)

    def test_deck_has_a_slide_per_page_before_the_figures(self):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        def pages():
            out = []
            for pg in ip.plan([site_doc(2)], per_sheet=3):
                fig = Figure(figsize=pptx_export.FIGURE_SIZE, dpi=60)
                pg.draw(fig, (0, 0, 1, 1))
                buf = io.BytesIO()
                fig.savefig(buf, format="png")
                out.append({"title": pg.title, "png": buf.getvalue(),
                            "notes": pg.notes()})
            return out

        def fig_png(number, fig):
            f = Figure(figsize=pptx_export.FIGURE_SIZE, dpi=40)
            buf = io.BytesIO()
            f.savefig(buf, format="png")
            return [buf.getvalue()]

        path = os.path.join(self.tmp, "d.pptx")
        figs = [{"name": "F", "caption": "", "state": {}}]
        n = pptx_export.build_deck(path, self.details, "", [], [], figs, fig_png,
                                   ("title", "images", "figures"), pages)
        prs = Presentation(path)
        slides = list(prs.slides)
        self.assertEqual(n, len(slides))
        titles = [s.shapes.title.text for s in slides]
        self.assertEqual(titles[-3:], ["Camera pictures", "SnapMap – S1",
                                       "Figure 1 – F"])
        cam = slides[-3]
        self.assertIn(MSO_SHAPE_TYPE.PICTURE, [sh.shape_type for sh in cam.shapes])
        self.assertIn("stage 10.000, 20.000 mm",
                      cam.notes_slide.notes_text_frame.text)
        # nothing to add: no callable, no slides
        n2 = pptx_export.build_deck(path, self.details, "", [], [], figs, fig_png,
                                    ("title", "images", "figures"), None)
        self.assertEqual(n2, n - 2)

    def test_the_deck_default_sections_include_images(self):
        self.assertEqual(pptx_export.SECTIONS,
                         ("title", "files", "metadata", "images", "figures"))


if __name__ == "__main__":
    unittest.main()
