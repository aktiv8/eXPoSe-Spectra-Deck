"""Repeat groups in the PDF, the slides and the Word report (they share
``Results.repeat_groups`` and ``resultspages.repeat_table``).

Run:  python -m unittest discover tests
"""

import os
import shutil
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
sys.path.insert(0, HERE)

import reportspec as rs  # noqa: E402
import resultspages as rp  # noqa: E402
from test_results import (HAVE_PDF, HAVE_PPTX, level, mupdf, row,  # noqa: E402
                          sample)

try:
    import docx  # noqa: F401
    HAVE_DOCX = True
except ImportError:
    HAVE_DOCX = False


def film(label, ti, states=True):
    """A one-level sample; ``ti`` scales Ti 2p against the fixed O 1s."""
    st = [("Ti-C", ti * .6), ("Ti-O", ti * .4)] if states else []
    lv = level(None, [row("Ti 2p", 2.0, ti, st), row("O 1s", 1.0, 100.0)])
    return sample(label, [lv])


def make(n=3, spec=True, extra=()):
    samples = [film(f"Film - spot {i + 1}", 100.0 + 10 * i) for i in range(n)]
    samples += list(extra)
    res = rp.Results(samples=samples)
    if spec:
        res.repeat_spec = {"Film": {"kind": "locations",
                                    "samples": [s.key for s in samples[:n]]}}
    return res


class TestResultsGroups(unittest.TestCase):
    def test_a_group_with_two_members_is_shown(self):
        (g,) = make().repeat_groups()
        self.assertEqual(g.name, "Film")
        self.assertEqual(len(g.members), 3)
        self.assertTrue(g.rows)

    def test_no_groups_when_none_were_made(self):
        self.assertEqual(make(spec=False).repeat_groups(), [])

    def test_a_sample_left_out_of_the_report_leaves_its_group(self):
        res = make()
        (g,) = res.repeat_groups(skip=[res.samples[2].key])
        self.assertEqual(len(g.members), 2)
        self.assertEqual(res.repeat_groups(
            skip=[res.samples[1].key, res.samples[2].key]), [])

    def test_a_missing_key_is_ignored_without_a_note(self):
        res = make()
        res.repeat_spec["Film"]["samples"].append("f1/gone")
        (g,) = res.repeat_groups()
        self.assertFalse(any("gone" in n for n in g.notes))

    def test_table_columns(self):
        (g,) = make().repeat_groups()
        title, header, rows, weights, left_out = rp.repeat_table(g, 180)
        self.assertEqual(header[:5], ["Region", "State", "Mean at %",
                                      "SD (n-1)", "n"])
        self.assertEqual(len(header), 8)
        self.assertAlmostEqual(sum(weights), 180)
        self.assertEqual(left_out, "")
        self.assertIn("locations", title.replace("different ", ""))

    def test_many_members_are_left_to_the_csv(self):
        res = make(n=6)
        (g,) = res.repeat_groups()
        _t, header, rows, _w, left_out = rp.repeat_table(g)
        self.assertEqual(len(header), 5)
        self.assertIn("repeat-group CSV", left_out)
        self.assertTrue(all(len(c) == 5 for _k, c in rows))


@unittest.skipUnless(HAVE_PDF, "reportlab / PyMuPDF / matplotlib missing")
class TestPdf(unittest.TestCase):
    def build(self, results, spec=None):
        import report
        from test_metasummary import Doc, region
        d = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, d, True)
        path = os.path.join(d, "r.pdf")
        report.build_report(
            path, {"title": "T", "summary": "Fine.", "methods": "Recorded."},
            "", [], [Doc([region("Survey", 160, step=1.0)], "a.vgd")], [],
            lambda *a: 0, spec=spec or rs.default_spec(), results=results)
        doc = mupdf.open(path)
        self.addCleanup(doc.close)
        return doc, "\n".join(p.get_text() for p in doc)

    def test_the_groups_follow_the_samples_with_their_numbers(self):
        doc, text = self.build(make())
        for token in ("Repeat groups", "Film: different locations",
                      "Mean at %", "SD (n-1)", "Ti 2p", "O 1s", "Ti-C",
                      "spot 1", "standard deviation (n"):
            self.assertIn(token, text)
        self.assertLess(text.index("Film - spot 3"), text.index(
            "Repeat groups"))
        self.assertIn("Repeat groups", [t for _l, t, _p in doc.get_toc()])

    def test_nothing_without_groups(self):
        _doc, text = self.build(make(spec=False))
        self.assertNotIn("Repeat groups", text)


@unittest.skipUnless(HAVE_PPTX and HAVE_PDF, "python-pptx not installed")
class TestDeck(unittest.TestCase):
    def build(self, results):
        import pptx_export
        from pptx import Presentation
        d = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, d, True)
        path = os.path.join(d, "d.pptx")
        pptx_export.build_deck(path, {"title": "T", "summary": "Fine."}, "",
                               [], [], [], lambda n, f: [],
                               spec=rs.default_spec(), results=results)
        return list(Presentation(path).slides)

    @staticmethod
    def title(slide):
        return slide.shapes.title.text if slide.shapes.title is not None \
            else ""

    def test_a_slide_per_group_after_the_samples(self):
        slides = self.build(make())
        titles = [self.title(s) for s in slides]
        mine = "Quantification – repeat group: Film"
        self.assertIn(mine, titles)
        self.assertGreater(titles.index(mine), titles.index(
            "Quantification – Film - spot 3"))
        table = [sh for sh in slides[titles.index(mine)].shapes
                 if sh.has_table][0].table
        self.assertEqual(table.cell(0, 2).text, "Mean at %")
        self.assertEqual(table.cell(1, 0).text, "Ti 2p")
        self.assertEqual(table.cell(1, 4).text, "3")

    def test_no_slide_without_groups(self):
        titles = [self.title(s) for s in self.build(make(spec=False))]
        self.assertFalse(any("repeat group" in t for t in titles))


@unittest.skipUnless(HAVE_DOCX and HAVE_PDF, "python-docx not installed")
class TestDocx(unittest.TestCase):
    def build(self, results):
        import docx_export
        from docx import Document
        from test_metasummary import Doc, region
        d = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, d, True)
        path = os.path.join(d, "d.docx")
        docx_export.build_document(
            path, {"title": "T", "summary": "Fine.", "methods": "Rec."}, "",
            [], [Doc([region("Survey", 160, step=1.0)], "a.vgd")], [],
            lambda n, f: [], spec=rs.default_spec(), results=results)
        return Document(path)

    def test_heading_and_table(self):
        doc = self.build(make())
        heads = [p.text for p in doc.paragraphs
                 if p.style.name.startswith("Heading")]
        self.assertIn("Repeat groups", heads)
        tables = [t for t in doc.tables
                  if len(t.rows[0].cells) > 2
                  and t.rows[0].cells[2].text == "Mean at %"]
        self.assertEqual(len(tables), 1)
        self.assertEqual(tables[0].rows[1].cells[0].text, "Ti 2p")

    def test_no_heading_without_groups(self):
        doc = self.build(make(spec=False))
        self.assertNotIn("Repeat groups", [p.text for p in doc.paragraphs])


if __name__ == "__main__":
    unittest.main()
